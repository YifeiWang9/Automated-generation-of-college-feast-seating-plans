import argparse
import csv
import re
from collections import Counter, defaultdict
from dataclasses import dataclass
from itertools import combinations
from pathlib import Path
from typing import Optional

import openpyxl


WORKBOOK = Path(
    "/Users/yifeiwang/Desktop/dissertation/data/"
    "MSc project workbook.xlsx"
)

CSV_DIR = Path(
    "/Users/yifeiwang/Desktop/dissertation/code/csv data"
)

SEATS_CSV = CSV_DIR / "all_layouts_seats.csv"
PAIRS_CSV = CSV_DIR / "all_layouts_seat_pairs.csv"

OUTPUT_DIR = Path("/Users/yifeiwang/Desktop/dissertation/code")
STAGE1_OUTPUT_CSV = OUTPUT_DIR / "college_feast_stage1_assignments.csv"
MAIN_OUTPUT_CSV = OUTPUT_DIR / "college_feast_main_course_solution.csv"
DESSERT_OUTPUT_CSV = OUTPUT_DIR / "college_feast_dessert_solution.csv"

SHEET = "College Feast"
MAIN_LAYOUT = "hall_110"
DESSERT_LAYOUTS = ("combination_60", "upper_45")

LAMBDA_1_GENDER = 1.0
LAMBDA_2_FELLOW_SPREAD = 1.0
LAMBDA_3_DESSERT_GUEST_MIX = 1.0
LAMBDA_4_REPEATED_CLOSENESS = 1.0
LAMBDA_5_PRESTIGE = 1.0
LAMBDA_6_DESSERT_TABLE_GENDER = 1.0
LAMBDA_7_DESSERT_TABLE_CG_MIX = 1.0
LAMBDA_8_STAGE1_HIGH_TABLE_SENIORITY = 1.0

STAGE_1_TIME_LIMIT = 900
STAGE_2_TIME_LIMIT = 3600
STAGE_3_TIME_LIMIT = 3600
MIP_RELATIVE_GAP = 0.01

SECTION_CODES = {
    "COLLEGE GUEST": "CG",
    "FELLOWS": "F",
    "BYE-FELLOWS": "BF",
    "COLLEGE LECTURERS": "CL",
    "HONORARY FELLOWS": "HF",
    "EMERITUS FELLOWS": "EF",
}

TITLE_WORDS = {
    "baroness", "dame", "dr", "hon", "lady", "lord", "mgr",
    "miss", "mr", "mrs", "ms", "prof", "professor", "qc",
    "revd", "reverend", "rt", "sir", "the",
}

FELLOW_SPREAD_CATEGORIES = {
    "FELLOWS",
    "BYE-FELLOWS",
    "HONORARY FELLOWS",
    "EMERITUS FELLOWS",
}

RANKED_CATEGORIES = {
    "FELLOWS",
    "HONORARY FELLOWS",
    "EMERITUS FELLOWS",
}


@dataclass(frozen=True)
class Diner:
    id: str
    excel_row: int
    category: str
    role: str
    no: int
    title: str
    first: str
    surname: str
    name: str
    gender: str
    notes: str
    rank_order: Optional[int]
    fellow_id: Optional[str] = None


def clean(value):
    """
    cleans the value so blanks and extra spaces do not cause problems later.
    """
    if value is None:
        return ""

    return re.sub(r"\s+", " ", str(value)).strip()


def slug(value):
    """
    turns text into a simple lower case label for ids and variable names.
    """
    text = clean(value).lower()
    text = re.sub(r"[^a-z0-9]+", "_", text).strip("_")

    return text or "x"


def to_int_or_none(value):
    """
    tries to change the value into an integer, and returns none if it cannot.
    """
    if value is None:
        return None

    if isinstance(value, int):
        return value

    if isinstance(value, float) and value.is_integer():
        return int(value)

    text = clean(value)

    if text.isdigit():
        return int(text)

    return None


def ordered_pair(diner_i, diner_j):
    """
    puts two diner ids into the same order every time, so pairs are easier to compare.
    """
    return tuple(sorted((diner_i, diner_j)))


def determine_gender(title_or_name, gender_value):
    """
    works out the gender code from the title or the gender column.
    """
    text = clean(title_or_name).lower().replace(".", "")
    explicit_gender = clean(gender_value).lower()
    words = set(re.findall(r"[a-z]+", text))

    if explicit_gender in {"m", "male", "man"}:
        return "M"

    if words.intersection({"mr", "sir", "lord", "mgr"}):
        return "M"

    if words.intersection(
        {"mrs", "miss", "ms", "madam", "dame", "lady", "baroness"}
    ):
        return "W"

    # In this workbook a blank gender cell means female.
    if explicit_gender in {"", "f", "female", "w", "woman"}:
        return "W"

    raise ValueError(
        "Cannot determine gender: "
        f"title_or_name={title_or_name!r}, gender={gender_value!r}"
    )


def person_name(title, first, surname):
    """
    joins the title, first name, and surname into one full name.
    """
    return " ".join(
        part
        for part in (clean(title), clean(first), clean(surname))
        if part
    )


def tokenize_name(text):
    """
    splits a name into useful words and removes title words.
    """
    tokens = re.findall(r"[a-z]+", clean(text).lower())
    return {
        token
        for token in tokens
        if token not in TITLE_WORDS
    }


def load_diners():
    """
    loads all valid diners from the workbook, including fellow guests when they exist.
    """
    workbook = openpyxl.load_workbook(
        WORKBOOK,
        data_only=True,
    )

    worksheet = workbook[SHEET]

    diners = []
    fellow_guest_pairs = []
    current_section = None
    fellow_guest_count = 0
    section_names = set(SECTION_CODES)

    for excel_row in range(4, worksheet.max_row + 1):
        values = [
            worksheet.cell(excel_row, column).value
            for column in range(1, 10)
        ]

        (
            number,
            title,
            gender,
            first,
            surname,
            guest_gender,
            guest_name,
            dinner,
            notes,
        ) = values

        detected_section = None

        for value in values[:4]:
            value_text = clean(value).upper()

            if value_text in section_names:
                detected_section = value_text
                break

        if detected_section is not None:
            current_section = detected_section
            continue

        if current_section is None:
            continue

        if (
            number is None
            and first is None
            and surname is None
        ):
            continue

        if to_int_or_none(dinner) is None:
            continue

        number_integer = to_int_or_none(number)

        if number_integer is None:
            number_integer = excel_row

        category_code = SECTION_CODES[current_section]
        diner_id = (
            f"{category_code}"
            f"{number_integer:03d}_"
            f"{slug(first)}_{slug(surname)}"
        )

        role = (
            "COLLEGE_GUEST"
            if current_section == "COLLEGE GUEST"
            else "COLLEGE_MEMBER"
        )

        rank_order = (
            number_integer
            if current_section in RANKED_CATEGORIES
            else None
        )

        diners.append(
            Diner(
                id=diner_id,
                excel_row=excel_row,
                category=current_section,
                role=role,
                no=number_integer,
                title=clean(title),
                first=clean(first),
                surname=clean(surname),
                name=person_name(title, first, surname),
                gender=determine_gender(title, gender),
                notes=clean(notes),
                rank_order=rank_order,
            )
        )

        guest_name_text = clean(guest_name)

        if guest_name_text and guest_name_text.lower() != "x":
            fellow_guest_count += 1
            guest_id = (
                f"FG{fellow_guest_count:03d}_"
                f"{slug(guest_name_text)}"
            )

            diners.append(
                Diner(
                    id=guest_id,
                    excel_row=excel_row,
                    category="FELLOW_GUEST",
                    role="FELLOW_GUEST",
                    no=fellow_guest_count,
                    title="",
                    first=guest_name_text,
                    surname="",
                    name=guest_name_text,
                    gender=determine_gender(guest_name_text, guest_gender),
                    notes="",
                    rank_order=None,
                    fellow_id=diner_id,
                )
            )

            fellow_guest_pairs.append((diner_id, guest_id))

    category_counts = Counter(diner.category for diner in diners)

    missing_sections = set(SECTION_CODES) - set(category_counts)

    if missing_sections:
        raise ValueError(
            "These College Feast sections were not detected: "
            f"{sorted(missing_sections)}"
        )

    return diners, fellow_guest_pairs


def main_seat_key(layout, seat_id):
    if layout != MAIN_LAYOUT:
        raise ValueError(f"Expected {MAIN_LAYOUT}, received {layout}.")

    return seat_id


def dessert_seat_key(layout, seat_id):
    return f"{layout}:{seat_id}"


def load_layouts():
    """
    loads the main hall and dessert seat layouts, including high table seats and close pairs.
    """
    seats_by_key = {}
    main_seats = []
    dessert_seats = []
    main_sections = defaultdict(list)
    dessert_rooms = defaultdict(list)
    dessert_table_sections = defaultdict(list)
    high_table = set()
    high_end = set()

    with SEATS_CSV.open(
        encoding="utf-8-sig",
        newline="",
    ) as file:
        reader = csv.DictReader(file)

        for row in reader:
            layout = row["layout"]

            if layout == MAIN_LAYOUT:
                seat_key = main_seat_key(layout, row["seat_id"])
            elif layout in DESSERT_LAYOUTS:
                seat_key = dessert_seat_key(layout, row["seat_id"])
            else:
                continue

            seat_record = {
                "key": seat_key,
                "seat_id": row["seat_id"],
                "layout": layout,
                "section": row["section_id"],
                "table": row["table_id"],
                "side": row["side"],
                "position": int(row["position"]),
                "excel_row": int(row["excel_row"]),
                "excel_col": int(row["excel_col"]),
                "is_high_table": clean(row["is_high_table"]) == "1",
                "high_table_zone": clean(row["high_table_zone"]),
            }

            seats_by_key[seat_key] = seat_record

            if layout == MAIN_LAYOUT:
                main_seats.append(seat_key)
                main_sections[row["section_id"]].append(seat_key)

                if seat_record["is_high_table"]:
                    high_table.add(seat_key)

                if seat_record["high_table_zone"].upper() == "END":
                    high_end.add(seat_key)
            else:
                dessert_seats.append(seat_key)
                dessert_rooms[layout].append(seat_key)
                dessert_table_sections[
                    f"{layout}:{row['section_id']}"
                ].append(seat_key)

    if not main_seats:
        raise ValueError("No Hall seats were loaded.")

    if not dessert_seats:
        raise ValueError("No dessert seats were loaded.")

    master_seat = "I8"

    if master_seat not in high_table:
        raise ValueError(
            f"Master seat {master_seat} is not a High Table seat."
        )

    high_centre = {
        master_seat,
    }

    high_vip = {
        "H8",
        "K8",
        "H12",
        "I12",
        "J12",
        "K12",
    }

    missing_vip = high_vip - high_table

    if missing_vip:
        raise ValueError(
            "These VIP seats are not High Table seats: "
            f"{sorted(missing_vip)}"
        )

    if master_seat in high_vip:
        raise ValueError("The Master's seat must not be in high_vip.")

    close_weights = {
        1: {},
        2: {},
    }

    with PAIRS_CSV.open(
        encoding="utf-8-sig",
        newline="",
    ) as file:
        reader = csv.DictReader(file)

        for row in reader:
            layout = row["layout"]
            weight = float(row["weight"])

            if layout == MAIN_LAYOUT:
                seat_s = main_seat_key(layout, row["seat_s"])
                seat_r = main_seat_key(layout, row["seat_r"])
                close_weights[1][seat_s, seat_r] = weight
            elif layout in DESSERT_LAYOUTS:
                seat_s = dessert_seat_key(layout, row["seat_s"])
                seat_r = dessert_seat_key(layout, row["seat_r"])
                close_weights[2][seat_s, seat_r] = weight

    return {
        "seats_by_key": seats_by_key,
        "main_seats": main_seats,
        "dessert_seats": dessert_seats,
        "main_sections": dict(main_sections),
        "dessert_rooms": dict(dessert_rooms),
        "dessert_table_sections": dict(dessert_table_sections),
        "close_weights": close_weights,
        "high_table": high_table,
        "high_vip": high_vip,
        "high_centre": high_centre,
        "high_end": high_end,
        "master_seat": master_seat,
    }


def build_neighbourhoods(seats, close_weights):
    close = defaultdict(list)
    near = defaultdict(list)
    with_area = defaultdict(list)

    for seat in seats:
        close[seat] = []
        near[seat] = []
        with_area[seat] = []

    for (seat_s, seat_r), weight in close_weights.items():
        close[seat_s].append(seat_r)
        near[seat_s].append(seat_r)

        if weight >= 0.999:
            with_area[seat_s].append(seat_r)

    return close, near, with_area


def derive_sets(diners, fellow_guest_pairs):
    """
    creates the main diner groups that are used later in the optimisation model.
    """
    diner_ids = [diner.id for diner in diners]
    diner_by_id = {diner.id: diner for diner in diners}

    men = {
        diner.id
        for diner in diners
        if diner.gender == "M"
    }

    women = {
        diner.id
        for diner in diners
        if diner.gender == "W"
    }

    college_guests = {
        diner.id
        for diner in diners
        if diner.role == "COLLEGE_GUEST"
    }

    named_fellow_guests = {
        diner.id
        for diner in diners
        if diner.role == "FELLOW_GUEST"
    }

    fellow_spread = {
        diner.id
        for diner in diners
        if diner.category in FELLOW_SPREAD_CATEGORIES
    }

    ranked = {
        diner.id
        for diner in diners
        if diner.category in RANKED_CATEGORIES
    }

    category_sizes = Counter(
        diner.category
        for diner in diners
        if diner.category in RANKED_CATEGORIES
    )

    importance = {
        diner.id: (
            category_sizes[diner.category] - diner.rank_order + 1
        ) / category_sizes[diner.category]
        for diner in diners
        if diner.rank_order is not None
    }

    high_table_diners = set()
    high_vip_diners = set()
    high_centre_diners = set()
    high_end_diners = set()

    for diner in diners:
        note = diner.notes.lower()

        if "high table" in note:
            high_table_diners.add(diner.id)

        if "vip" in note and "high" in note:
            high_vip_diners.add(diner.id)
            high_table_diners.add(diner.id)

        if (
            ("centre" in note or "center" in note)
            and "high" in note
        ):
            high_centre_diners.add(diner.id)
            high_table_diners.add(diner.id)

        if "end" in note and "high" in note:
            high_end_diners.add(diner.id)
            high_table_diners.add(diner.id)

    master_candidates = [
        diner.id
        for diner in diners
        if "master" in (
            f"{diner.title} {diner.first} "
            f"{diner.surname} {diner.notes}"
        ).lower()
    ]

    if len(master_candidates) != 1:
        raise ValueError(
            "Exactly one Master must be identified. "
            f"Found: {master_candidates}"
        )

    master_diner = master_candidates[0]

    if master_diner not in high_centre_diners:
        high_centre_diners.add(master_diner)
        high_table_diners.add(master_diner)

    pair_set = {
        ordered_pair(fellow_id, guest_id)
        for fellow_id, guest_id in fellow_guest_pairs
    }

    mix_pairs = []

    for diner_i, diner_j in combinations(diner_ids, 2):
        pair = ordered_pair(diner_i, diner_j)

        if pair in pair_set:
            continue

        if (
            diner_i in college_guests
            or diner_j in college_guests
            or diner_i in named_fellow_guests
            or diner_j in named_fellow_guests
        ):
            mix_pairs.append(pair)

    return {
        "I": diner_ids,
        "diner_by_id": diner_by_id,
        "M": men,
        "W": women,
        "I_CG": college_guests,
        "I_FG": named_fellow_guests,
        "F_spread": fellow_spread,
        "I_rank": ranked,
        "pi": importance,
        "I_H": high_table_diners,
        "I_VIP": high_vip_diners,
        "I_centre": high_centre_diners,
        "I_end": high_end_diners,
        "master_diner": master_diner,
        "P_FG": fellow_guest_pairs,
        "P_FG_unordered": pair_set,
        "P_mix": mix_pairs,
    }


def match_diner_reference(reference_text, diners):
    target_tokens = tokenize_name(reference_text)

    if not target_tokens:
        return None

    matches = []

    for diner in diners:
        diner_tokens = tokenize_name(diner.name)

        if target_tokens.issubset(diner_tokens):
            matches.append(diner.id)

    if len(matches) == 1:
        return matches[0]

    return None


def derive_conflicts(diners, sets):
    """
    reads the notes and turns do-not-seat rules into conflict pairs.
    """
    individual_near = set()
    individual_with = set()
    no_with_college_guests = set()
    unmatched_notes = []

    for diner in diners:
        note = diner.notes.lower()

        if "do not" not in note:
            continue

        if "college guest" in note:
            no_with_college_guests.add(diner.id)
            continue

        near_match = re.search(
            r"do not\s+(?:seat|sit|sear)\s+near\s+(.+)",
            note,
        )

        with_match = re.search(
            r"do not\s+(?:seat|sit)\s+with\s+(.+)",
            note,
        )

        if near_match:
            target = match_diner_reference(
                near_match.group(1),
                diners,
            )

            if target is None:
                unmatched_notes.append((diner.id, diner.notes))
            else:
                individual_near.add(ordered_pair(diner.id, target))

            continue

        if with_match:
            target = match_diner_reference(
                with_match.group(1),
                diners,
            )

            if target is None:
                unmatched_notes.append((diner.id, diner.notes))
            else:
                individual_with.add(ordered_pair(diner.id, target))

            continue

        unmatched_notes.append((diner.id, diner.notes))

    college_guest_with = {
        ordered_pair(diner_id, college_guest_id)
        for diner_id in no_with_college_guests
        for college_guest_id in sets["I_CG"]
        if diner_id != college_guest_id
    }

    all_with = individual_with.union(college_guest_with)

    return {
        "E_near": individual_near,
        "E_individual_with": individual_with,
        "E_CG_with": college_guest_with,
        "E_with": all_with,
        "I_NWCG": no_with_college_guests,
        "unmatched_notes": unmatched_notes,
    }


def build_stage_sets(layouts):
    main_seats = layouts["main_seats"]
    dessert_seats = layouts["dessert_seats"]
    main_sections = layouts["main_sections"]
    dessert_rooms = layouts["dessert_rooms"]
    dessert_table_sections = layouts["dessert_table_sections"]
    seats_by_key = layouts["seats_by_key"]
    close_weights = layouts["close_weights"]

    main_close, main_near, main_with = build_neighbourhoods(
        main_seats,
        close_weights[1],
    )

    dessert_close, dessert_near, dessert_with = build_neighbourhoods(
        dessert_seats,
        close_weights[2],
    )

    dessert_fg = defaultdict(list)

    for seat in dessert_seats:
        dessert_fg[seat] = []

    for seat_s, seat_r in close_weights[2]:
        if seat_r in dessert_with[seat_s]:
            continue

        if (
            seats_by_key[seat_s]["layout"]
            == seats_by_key[seat_r]["layout"]
        ):
            dessert_fg[seat_s].append(seat_r)

    gamma_main_with = {
        section: set()
        for section in main_sections
    }

    for seat_s in main_seats:
        section_s = seats_by_key[seat_s]["section"]

        for seat_r in main_with[seat_s]:
            section_r = seats_by_key[seat_r]["section"]
            gamma_main_with[section_s].add(section_r)

    return {
        "S1": main_seats,
        "S2": dessert_seats,
        "Q1": list(main_sections),
        "Q2": list(dessert_rooms),
        "T2": list(dessert_table_sections),
        "S_q1": main_sections,
        "S_q2": dessert_rooms,
        "S_t2": dessert_table_sections,
        "N1": main_close,
        "N1_near": main_near,
        "N1_with": main_with,
        "N2": dessert_close,
        "N2_near": dessert_near,
        "N2_with": dessert_with,
        "N2_FG": dessert_fg,
        "Gamma1_with": gamma_main_with,
    }


def seat_prestige(layouts):
    prestige = {}

    for seat in layouts["main_seats"]:
        if seat in layouts["high_centre"]:
            prestige[seat] = 10.0
        elif seat in layouts["high_vip"]:
            prestige[seat] = 8.0
        elif seat in layouts["high_table"] - layouts["high_end"]:
            prestige[seat] = 5.0
        elif seat in layouts["high_end"]:
            prestige[seat] = 3.0
        else:
            prestige[seat] = 0.0

    return prestige


def validate_data(diners, layouts, sets, conflicts, stage_sets):
    if len(sets["I"]) > len(stage_sets["S1"]):
        raise ValueError("There are more diners than Hall seats.")

    if len(sets["I"]) > len(stage_sets["S2"]):
        raise ValueError("There are more diners than dessert seats.")

    if len(sets["I_H"]) > len(layouts["high_table"]):
        raise ValueError(
            "There are more required High Table diners "
            "than High Table seats."
        )

    if not layouts["high_vip"].issubset(layouts["high_table"]):
        raise ValueError("high_vip must be a subset of High Table seats.")

    if not layouts["high_centre"].issubset(layouts["high_table"]):
        raise ValueError("high_centre must be a subset of High Table seats.")

    if layouts["master_seat"] in layouts["high_vip"]:
        raise ValueError("The Master's seat cannot be a VIP seat.")

    for fellow_id, guest_id in sets["P_FG"]:
        if fellow_id not in sets["I"] or guest_id not in sets["I"]:
            raise ValueError(f"Unknown Fellow-Guest pair: {(fellow_id, guest_id)}")

    for note in conflicts["unmatched_notes"]:
        print("WARNING unmatched note:", note)

    print("\nDATA SUMMARY")
    print("Diners |I|:", len(sets["I"]))
    print("Hall seats |S1|:", len(stage_sets["S1"]))
    print("Dessert seats |S2|:", len(stage_sets["S2"]))
    print("Main sections |Q1|:", len(stage_sets["Q1"]), stage_sets["Q1"])
    print("Dessert rooms |Q2|:", len(stage_sets["Q2"]), stage_sets["Q2"])
    print("Dessert table sections |T2|:", len(stage_sets["T2"]))
    print("Fellow-Guest pairs |P_FG|:", len(sets["P_FG"]))
    print("College Guests |I_CG|:", len(sets["I_CG"]))
    print("Named Fellow Guests |I_FG|:", len(sets["I_FG"]))
    print("Fellow spread diners |F_spread|:", len(sets["F_spread"]))
    print("High Table seats |H|:", len(layouts["high_table"]))
    print("High Table diners |I_H|:", len(sets["I_H"]))
    print("VIP High Table diners |I_VIP|:", len(sets["I_VIP"]))
    print("Centre High Table diners |I_centre|:", len(sets["I_centre"]))
    print("End High Table diners |I_end|:", len(sets["I_end"]))
    print("Master diner:", sets["master_diner"])
    print("Master seat:", layouts["master_seat"])
    print("Gender counts:", Counter(diner.gender for diner in diners))
    print("Category counts:", Counter(diner.category for diner in diners))
    print("Do-not-near pairs |E_near|:", len(conflicts["E_near"]))
    print("Do-not-with pairs |E_with|:", len(conflicts["E_with"]))
    print(
        "Do-not-with College Guest generated pairs |E_CG_with|:",
        len(conflicts["E_CG_with"]),
    )


def import_xpress():
    import xpress as xp

    xpress_license = Path(
        "/Applications/FICO Xpress/"
        "xpressmp/bin/xpauth.xpr"
    )

    if xpress_license.exists():
        xp.init(str(xpress_license))

    return xp


def set_common_controls(problem, time_limit):
    problem.controls.outputlog = 1
    problem.controls.miprelstop = MIP_RELATIVE_GAP

    try:
        problem.controls.timelimit = time_limit
    except AttributeError:
        problem.controls.maxtime = time_limit


def require_solution(problem, stage_name):
    print(f"{stage_name} solve status:", problem.attributes.solvestatus)
    print(f"{stage_name} solution status:", problem.attributes.solstatus)
    print(f"{stage_name} number of solutions:", problem.attributes.mipsols)

    if problem.attributes.mipsols < 1:
        raise RuntimeError(f"{stage_name} did not find an integer solution.")

    print(f"{stage_name} objective:", problem.attributes.objval)


def xsum(xp, terms):
    terms = list(terms)

    if not terms:
        return 0

    return xp.Sum(terms)


def solve_stage_1(xp, sets, layouts, stage_sets):
    """
    solves stage 1 by assigning each diner to a main course section and a dessert room.
    """
    problem = xp.problem("college_feast_stage_1_sections_and_rooms")
    set_common_controls(problem, STAGE_1_TIME_LIMIT)

    I = sets["I"]
    M = sets["M"]
    W = sets["W"]
    F_spread = sets["F_spread"]
    I_CG = sets["I_CG"]
    I_FG = sets["I_FG"]
    Q1 = stage_sets["Q1"]
    Q2 = stage_sets["Q2"]
    S1 = stage_sets["S1"]
    S2 = stage_sets["S2"]
    S_q1 = stage_sets["S_q1"]
    S_q2 = stage_sets["S_q2"]
    q_high = layouts["seats_by_key"][layouts["master_seat"]]["section"]

    a_main = {}
    a_dessert = {}


    for diner_id in I:
        for section in Q1:
            a_main[diner_id, section] = xp.var(
                vartype=xp.binary,
                name=f"a_main_{diner_id}_{slug(section)}",
            )

        for room in Q2:
            a_dessert[diner_id, room] = xp.var(
                vartype=xp.binary,
                name=f"a_dessert_{diner_id}_{slug(room)}",
            )

    problem.addVariable(
        list(a_main.values())
        + list(a_dessert.values())
    )

    d_gender_plus = {}
    d_gender_minus = {}

    for arrangement, sections in ((1, Q1), (2, Q2)):
        for group in ("M", "W"):
            for section in sections:
                d_gender_plus[arrangement, group, section] = xp.var(
                    lb=0,
                    name=(
                        f"dplus_gender_{arrangement}_{group}_"
                        f"{slug(section)}"
                    ),
                )
                d_gender_minus[arrangement, group, section] = xp.var(
                    lb=0,
                    name=(
                        f"dminus_gender_{arrangement}_{group}_"
                        f"{slug(section)}"
                    ),
                )

    d_fellow_plus = {}
    d_fellow_minus = {}

    for section in Q1:
        d_fellow_plus[section] = xp.var(
            lb=0,
            name=f"dplus_fellow_spread_{slug(section)}",
        )
        d_fellow_minus[section] = xp.var(
            lb=0,
            name=f"dminus_fellow_spread_{slug(section)}",
        )

    d_cg_plus = {}
    d_cg_minus = {}
    d_fg_plus = {}
    d_fg_minus = {}

    for room in Q2:
        d_cg_plus[room] = xp.var(
            lb=0,
            name=f"dplus_cg_{slug(room)}",
        )
        d_cg_minus[room] = xp.var(
            lb=0,
            name=f"dminus_cg_{slug(room)}",
        )
        d_fg_plus[room] = xp.var(
            lb=0,
            name=f"dplus_fg_{slug(room)}",
        )
        d_fg_minus[room] = xp.var(
            lb=0,
            name=f"dminus_fg_{slug(room)}",
        )

    problem.addVariable(
        list(d_gender_plus.values())
        + list(d_gender_minus.values())
        + list(d_fellow_plus.values())
        + list(d_fellow_minus.values())
        + list(d_cg_plus.values())
        + list(d_cg_minus.values())
        + list(d_fg_plus.values())
        + list(d_fg_minus.values())
    )

    # each diner must be assigned to exactly one main course section and one dessert room.
    for diner_id in I:
        problem.addConstraint(
            xp.Sum(
                a_main[diner_id, section]
                for section in Q1
            ) == 1
        )
        problem.addConstraint(
            xp.Sum(
                a_dessert[diner_id, room]
                for room in Q2
            ) == 1
        )

    # each main course section can only take up to the number of seats it has.
    for section in Q1:
        problem.addConstraint(
            xp.Sum(
                a_main[diner_id, section]
                for diner_id in I
            ) <= len(S_q1[section])
        )

    # each dessert room can only take up to the number of seats it has.
    for room in Q2:
        problem.addConstraint(
            xp.Sum(
                a_dessert[diner_id, room]
                for diner_id in I
            ) <= len(S_q2[room])
        )

    # required high table diners must be assigned to the high table section.
    for diner_id in sets["I_H"]:
        problem.addConstraint(a_main[diner_id, q_high] == 1)

    # fellow and guest pairs need to be close in the main course and stay in the same dessert room.
    for fellow_id, guest_id in sets["P_FG"]:
        for section in Q1:
            compatible_sections = stage_sets["Gamma1_with"][section]
            problem.addConstraint(
                a_main[fellow_id, section]
                <= xsum(
                    xp,
                    (
                        a_main[guest_id, compatible_section]
                        for compatible_section in compatible_sections
                    ),
                )
            )

        for room in Q2:
            problem.addConstraint(
                a_dessert[fellow_id, room]
                == a_dessert[guest_id, room]
            )

    stage_groups = {
        "M": M,
        "W": W,
    }

    # these constraints measure how far each section or room is from the target gender balance.
    for group_name, group_set in stage_groups.items():
        for section in Q1:
            target = len(group_set) * len(S_q1[section]) / len(S1)
            actual = xp.Sum(
                a_main[diner_id, section]
                for diner_id in group_set
            )
            problem.addConstraint(
                actual - target
                == (
                    d_gender_plus[1, group_name, section]
                    - d_gender_minus[1, group_name, section]
                )
            )

        for room in Q2:
            target = len(group_set) * len(S_q2[room]) / len(S2)
            actual = xp.Sum(
                a_dessert[diner_id, room]
                for diner_id in group_set
            )
            problem.addConstraint(
                actual - target
                == (
                    d_gender_plus[2, group_name, room]
                    - d_gender_minus[2, group_name, room]
                )
            )

    # these constraints measure how evenly fellows are spread across main course sections.
    for section in Q1:
        target = len(F_spread) * len(S_q1[section]) / len(S1)
        actual = xp.Sum(
            a_main[diner_id, section]
            for diner_id in F_spread
        )
        problem.addConstraint(
            actual - target
            == d_fellow_plus[section] - d_fellow_minus[section]
        )

    # these constraints measure how evenly college guests and fellow guests are spread across dessert rooms.
    for room in Q2:
        target_cg = len(I_CG) * len(S_q2[room]) / len(S2)
        actual_cg = xp.Sum(
            a_dessert[diner_id, room]
            for diner_id in I_CG
        )
        problem.addConstraint(
            actual_cg - target_cg
            == d_cg_plus[room] - d_cg_minus[room]
        )

        target_fg = len(I_FG) * len(S_q2[room]) / len(S2)
        actual_fg = xp.Sum(
            a_dessert[diner_id, room]
            for diner_id in I_FG
        )
        problem.addConstraint(
            actual_fg - target_fg
            == d_fg_plus[room] - d_fg_minus[room]
        )


    # NORMALISED STAGE 1 OBJECTIVE
    gender_scale_stage1 = (
        4.0 * (len(M) + len(W))
    )

    fellow_scale_stage1 = (
        2.0 * len(F_spread)
    )

    dessert_guest_scale_stage1 = (
        2.0 * (len(I_CG) + len(I_FG))
    )

    gender_term_stage1 = (
        xp.Sum(
            d_gender_plus[1, group, section]
            + d_gender_minus[1, group, section]
            for group in ("M", "W")
            for section in Q1
        )
        + xp.Sum(
            d_gender_plus[2, group, room]
            + d_gender_minus[2, group, room]
            for group in ("M", "W")
            for room in Q2
        )
    ) / gender_scale_stage1

    fellow_term_stage1 = (
        xp.Sum(
            d_fellow_plus[section]
            + d_fellow_minus[section]
            for section in Q1
        )
        / fellow_scale_stage1
    )

    dessert_guest_term_stage1 = (
        xp.Sum(
            d_cg_plus[room]
            + d_cg_minus[room]
            + d_fg_plus[room]
            + d_fg_minus[room]
            for room in Q2
        )
        / dessert_guest_scale_stage1
    )

    high_table_seniority_scale_stage1 = sum(
        sets["pi"][diner_id]
        for diner_id in sets["I_rank"]
    )

    if high_table_seniority_scale_stage1 <= 0:
        raise ValueError(
            "Stage 1 High Table seniority scale must be positive."
        )

    high_table_seniority_term_stage1 = (
        xp.Sum(
            sets["pi"][diner_id]
            * a_main[diner_id, q_high]
            for diner_id in sets["I_rank"]
        )
        / high_table_seniority_scale_stage1
    )

    objective = (
        LAMBDA_1_GENDER
        * gender_term_stage1
        + LAMBDA_2_FELLOW_SPREAD
        * fellow_term_stage1
        + LAMBDA_3_DESSERT_GUEST_MIX
        * dessert_guest_term_stage1
        - LAMBDA_8_STAGE1_HIGH_TABLE_SENIORITY
        * high_table_seniority_term_stage1
    )



    problem.setObjective(objective, sense=xp.minimize)

    print("\nStage 1 model built. Solving...")
    problem.solve()
    require_solution(problem, "Stage 1")

    main_assignment = {}
    dessert_assignment = {}

    for diner_id in I:
        selected_main = [
            section
            for section in Q1
            if problem.getSolution(a_main[diner_id, section]) > 0.5
        ]
        selected_dessert = [
            room
            for room in Q2
            if problem.getSolution(a_dessert[diner_id, room]) > 0.5
        ]

        if len(selected_main) != 1:
            raise RuntimeError(
                f"{diner_id} does not have exactly one Main Course section."
            )

        if len(selected_dessert) != 1:
            raise RuntimeError(
                f"{diner_id} does not have exactly one dessert room."
            )

        main_assignment[diner_id] = selected_main[0]
        dessert_assignment[diner_id] = selected_dessert[0]

    return main_assignment, dessert_assignment


def solve_stage_2(xp, sets, conflicts, layouts, stage_sets, main_assignment):
    """
    solves stage 2 by assigning exact main course seats after stage 1 is fixed.
    """
    problem = xp.problem("college_feast_stage_2_main_course_seats")
    set_common_controls(problem, STAGE_2_TIME_LIMIT)

    I = sets["I"]
    S1 = stage_sets["S1"]
    Q1 = stage_sets["Q1"]
    S_q1 = stage_sets["S_q1"]
    N1_near = stage_sets["N1_near"]
    N1_with = stage_sets["N1_with"]

    x = {}


    for diner_id in I:

        for seat in S1:
            x[diner_id, seat] = xp.var(
                vartype=xp.binary,
                name=f"xM_{diner_id}_{slug(seat)}",
            )

    problem.addVariable(list(x.values()))

    # each diner must get exactly one main course seat.
    for diner_id in I:
        problem.addConstraint(
            xp.Sum(
                x[diner_id, seat]
                for seat in S1
            ) == 1
        )

    # each main course seat can have at most one diner.
    for seat in S1:
        problem.addConstraint(
            xp.Sum(
                x[diner_id, seat]
                for diner_id in I
            ) <= 1
        )

    # the exact main course seat must stay inside the section chosen in stage 1.
    for diner_id in I:
        for section in Q1:
            required = 1 if main_assignment[diner_id] == section else 0
            problem.addConstraint(
                xp.Sum(
                    x[diner_id, seat]
                    for seat in S_q1[section]
                ) == required
            )

    # high table diners must sit somewhere on the high table.
    for diner_id in sets["I_H"]:
        problem.addConstraint(
            xp.Sum(
                x[diner_id, seat]
                for seat in layouts["high_table"]
            ) == 1
        )

    # vip high table diners must sit in the vip high table seats.
    for diner_id in sets["I_VIP"]:
        problem.addConstraint(
            xp.Sum(
                x[diner_id, seat]
                for seat in layouts["high_vip"]
            ) == 1
        )

    # centre high table diners must sit in the centre high table seats.
    for diner_id in sets["I_centre"]:
        problem.addConstraint(
            xp.Sum(
                x[diner_id, seat]
                for seat in layouts["high_centre"]
            ) == 1
        )

    # end high table diners must sit in the end high table seats.
    for diner_id in sets["I_end"]:
        problem.addConstraint(
            xp.Sum(
                x[diner_id, seat]
                for seat in layouts["high_end"]
            ) == 1
        )

    # the master must sit in the fixed master seat.
    problem.addConstraint(
        x[sets["master_diner"], layouts["master_seat"]] == 1
    )

    # fellow and guest pairs must be seated next to or with each other at the main course.
    for fellow_id, guest_id in sets["P_FG"]:
        for seat in S1:
            problem.addConstraint(
                x[fellow_id, seat]
                <= xsum(
                    xp,
                    (
                        x[guest_id, neighbour]
                        for neighbour in N1_with[seat]
                    ),
                )
            )

    # diners with a do-not-near rule cannot be placed near each other.
    for diner_i, diner_j in conflicts["E_near"]:
        for seat in S1:
            problem.addConstraint(
                x[diner_i, seat]
                + xsum(
                    xp,
                    (
                        x[diner_j, neighbour]
                        for neighbour in N1_near[seat]
                    ),
                )
                <= 1
            )

    # diners with a do-not-with rule cannot be placed directly with each other.
    for diner_i, diner_j in conflicts["E_with"]:
        for seat in S1:
            problem.addConstraint(
                x[diner_i, seat]
                + xsum(
                    xp,
                    (
                        x[diner_j, neighbour]
                        for neighbour in N1_with[seat]
                    ),
                )
                <= 1
            )

    prestige = seat_prestige(layouts)

    prestige_scale = (
        max(prestige.values())
        * sum(
            sets["pi"][diner_id]
            for diner_id in sets["I_rank"]
        )
    )

    if prestige_scale <= 0:
        raise ValueError(
            "Prestige scale must be positive."
        )

    prestige_term_stage2 = (
        xp.Sum(
            sets["pi"][diner_id]
            * prestige[seat]
            * x[diner_id, seat]
            for diner_id in sets["I_rank"]
            for seat in S1
        )
        / prestige_scale
    )

    objective = (
        -LAMBDA_5_PRESTIGE
        * prestige_term_stage2
    )

    problem.setObjective(objective, sense=xp.minimize)

    print("\nStage 2 model built. Solving...")
    problem.solve()
    require_solution(problem, "Stage 2")

    main_seat = {}

    for diner_id in I:
        selected = [
            seat
            for seat in S1
            if problem.getSolution(x[diner_id, seat]) > 0.5
        ]

        if len(selected) != 1:
            raise RuntimeError(
                f"{diner_id} does not have exactly one Main Course seat."
            )

        main_seat[diner_id] = selected[0]

    return main_seat


def repeated_main_pairs(sets, layouts, main_seat):
    close_pairs = set(layouts["close_weights"][1])
    repeated_pairs = []

    for diner_i, diner_j in sets["P_mix"]:
        seat_i = main_seat[diner_i]
        seat_j = main_seat[diner_j]

        if (seat_i, seat_j) in close_pairs:
            repeated_pairs.append((diner_i, diner_j))

    return repeated_pairs


def solve_stage_3(
    xp,
    sets,
    conflicts,
    layouts,
    stage_sets,
    dessert_assignment,
    main_seat,
):
    """
    solves stage 3 by assigning dessert seats while avoiding repeated closeness.
    """
    problem = xp.problem("college_feast_stage_3_dessert_seats")
    set_common_controls(problem, STAGE_3_TIME_LIMIT)

    I = sets["I"]
    I_CG = sets["I_CG"]

    S2 = stage_sets["S2"]
    Q2 = stage_sets["Q2"]
    T2 = stage_sets["T2"]
    S_q2 = stage_sets["S_q2"]
    S_t2 = stage_sets["S_t2"]
    N2 = stage_sets["N2"]
    N2_near = stage_sets["N2_near"]
    N2_with = stage_sets["N2_with"]
    N2_FG = stage_sets["N2_FG"]

    x = {}

    for diner_id in I:
        for seat in S2:
            x[diner_id, seat] = xp.var(
                vartype=xp.binary,
                name=f"xD_{diner_id}_{slug(seat)}",
            )

    problem.addVariable(list(x.values()))

    d_gender_plus = {}
    d_gender_minus = {}

    for group in ("M", "W"):
        for table_section in T2:
            d_gender_plus[group, table_section] = xp.var(
                lb=0,
                name=f"dplus_dessert_{group}_{slug(table_section)}",
            )
            d_gender_minus[group, table_section] = xp.var(
                lb=0,
                name=f"dminus_dessert_{group}_{slug(table_section)}",
            )

    d_cg_table_plus = {}
    d_cg_table_minus = {}

    for table_section in T2:
        d_cg_table_plus[table_section] = xp.var(
            lb=0,
            name=f"dplus_dessert_cg_{slug(table_section)}",
        )
        d_cg_table_minus[table_section] = xp.var(
            lb=0,
            name=f"dminus_dessert_cg_{slug(table_section)}",
        )

    problem.addVariable(
        list(d_gender_plus.values())
        + list(d_gender_minus.values())
        + list(d_cg_table_plus.values())
        + list(d_cg_table_minus.values())
    )

    # each diner must get exactly one dessert seat.
    for diner_id in I:
        problem.addConstraint(
            xp.Sum(
                x[diner_id, seat]
                for seat in S2
            ) == 1
        )

    # each dessert seat can have at most one diner.
    for seat in S2:
        problem.addConstraint(
            xp.Sum(
                x[diner_id, seat]
                for diner_id in I
            ) <= 1
        )

    # the exact dessert seat must stay inside the dessert room chosen in stage 1.
    for diner_id in I:
        for room in Q2:
            required = 1 if dessert_assignment[diner_id] == room else 0
            problem.addConstraint(
                xp.Sum(
                    x[diner_id, seat]
                    for seat in S_q2[room]
                ) == required
            )

    # these constraints measure dessert table gender balance.
    for group_name, group_set in (("M", sets["M"]), ("W", sets["W"])):
        for table_section in T2:
            target = len(group_set) * len(S_t2[table_section]) / len(S2)
            actual = xp.Sum(
                x[diner_id, seat]
                for diner_id in group_set
                for seat in S_t2[table_section]
            )
            problem.addConstraint(
                actual - target
                == (
                    d_gender_plus[group_name, table_section]
                    - d_gender_minus[group_name, table_section]
                )
            )

    # these constraints measure how evenly college guests are spread across dessert tables.
    for table_section in T2:
        target_cg = (
            len(I_CG)
            * len(S_t2[table_section])
            / len(S2)
        )

        actual_cg = xp.Sum(
            x[diner_id, seat]
            for diner_id in I_CG
            for seat in S_t2[table_section]
        )

        problem.addConstraint(
            actual_cg - target_cg
            == (
                d_cg_table_plus[table_section]
                - d_cg_table_minus[table_section]
            )
        )

    # fellow and guest pairs should be close enough at dessert.
    for fellow_id, guest_id in sets["P_FG"]:
        for seat in S2:
            problem.addConstraint(
                x[fellow_id, seat]
                <= xsum(
                    xp,
                    (
                        x[guest_id, neighbour]
                        for neighbour in N2_FG[seat]
                    ),
                )
            )

    # do-not-near pairs are not allowed to sit near each other at dessert.
    for diner_i, diner_j in conflicts["E_near"]:
        for seat in S2:
            problem.addConstraint(
                x[diner_i, seat]
                + xsum(
                    xp,
                    (
                        x[diner_j, neighbour]
                        for neighbour in N2_near[seat]
                    ),
                )
                <= 1
            )

    # do-not-with pairs are not allowed to sit directly with each other at dessert.
    for diner_i, diner_j in conflicts["E_with"]:
        for seat in S2:
            problem.addConstraint(
                x[diner_i, seat]
                + xsum(
                    xp,
                    (
                        x[diner_j, neighbour]
                        for neighbour in N2_with[seat]
                    ),
                )
                <= 1
            )

    repeated_pairs = repeated_main_pairs(sets, layouts, main_seat)
    repeated = {}
    
    # create variables for pairs that might be close again at dessert.
    for diner_i, diner_j in repeated_pairs:
        repeated[diner_i, diner_j] = xp.var(
            vartype=xp.binary,
            name=f"rep_{diner_i}_{diner_j}",
        )

    if repeated:
        problem.addVariable(list(repeated.values()))
        
    # this marks pairs that were close in the main course and are close again at dessert.
    for diner_i, diner_j in repeated_pairs:
        rep_var = repeated[diner_i, diner_j]

        for seat in S2:
            problem.addConstraint(
                rep_var
                >= (
                    x[diner_i, seat]
                    + xsum(
                        xp,
                        (
                            x[diner_j, neighbour]
                            for neighbour in N2[seat]
                        ),
                    )
                    - 1
                )
            )


    # NORMALISED STAGE 3 OBJECTIVE
    if repeated_pairs:
        repeated_term_stage3 = (
            xsum(
                xp,
                (
                    repeated[pair]
                    for pair in repeated_pairs
                ),
            )
            / len(repeated_pairs)
        )
    else:
        repeated_term_stage3 = 0

    gender_scale_stage3 = (
        2.0 * (len(sets["M"]) + len(sets["W"]))
    )

    dessert_gender_term_stage3 = (
        xp.Sum(
            d_gender_plus[group, table_section]
            + d_gender_minus[group, table_section]
            for group in ("M", "W")
            for table_section in T2
        )
        / gender_scale_stage3
    )

    cg_scale_stage3 = (
        2.0 * len(I_CG)
    )

    dessert_cg_term_stage3 = (
        xp.Sum(
            d_cg_table_plus[table_section]
            + d_cg_table_minus[table_section]
            for table_section in T2
        )
        / cg_scale_stage3
    )

    objective = (
        LAMBDA_4_REPEATED_CLOSENESS
        * repeated_term_stage3
        + LAMBDA_6_DESSERT_TABLE_GENDER
        * dessert_gender_term_stage3
        + LAMBDA_7_DESSERT_TABLE_CG_MIX
        * dessert_cg_term_stage3
    )

    problem.setObjective(objective, sense=xp.minimize)

    print("\nStage 3 model built. Solving...")
    print("Main-course close P_mix pairs |R_M|:", len(repeated_pairs))
    problem.solve()
    require_solution(problem, "Stage 3")

    dessert_seat = {}

    for diner_id in I:
        selected = [
            seat
            for seat in S2
            if problem.getSolution(x[diner_id, seat]) > 0.5
        ]

        if len(selected) != 1:
            raise RuntimeError(
                f"{diner_id} does not have exactly one dessert seat."
            )

        dessert_seat[diner_id] = selected[0]

    return dessert_seat, repeated_pairs


def write_stage1_csv(sets, main_assignment, dessert_assignment):
    diner_by_id = sets["diner_by_id"]

    rows = []

    for diner_id in sets["I"]:
        diner = diner_by_id[diner_id]
        rows.append(
            {
                "diner_id": diner_id,
                "full_name": diner.name,
                "category": diner.category,
                "role": diner.role,
                "gender": diner.gender,
                "main_section": main_assignment[diner_id],
                "dessert_room": dessert_assignment[diner_id],
                "notes": diner.notes,
            }
        )

    rows.sort(key=lambda row: (row["main_section"], row["full_name"]))

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    with STAGE1_OUTPUT_CSV.open("w", encoding="utf-8", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def solution_rows(sets, layouts, assignment, arrangement):
    diner_by_id = sets["diner_by_id"]
    seats_by_key = layouts["seats_by_key"]
    rows = []

    for diner_id, seat_key in assignment.items():
        diner = diner_by_id[diner_id]
        seat = seats_by_key[seat_key]
        rows.append(
            {
                "arrangement": arrangement,
                "layout": seat["layout"],
                "seat_key": seat_key,
                "seat_id": seat["seat_id"],
                "seat_section": seat["section"],
                "seat_table": seat["table"],
                "seat_side": seat["side"],
                "seat_position": seat["position"],
                "excel_row": seat["excel_row"],
                "excel_col": seat["excel_col"],
                "diner_id": diner_id,
                "category": diner.category,
                "role": diner.role,
                "title": diner.title,
                "first": diner.first,
                "surname": diner.surname,
                "full_name": diner.name,
                "gender": diner.gender,
                "notes": diner.notes,
            }
        )

    rows.sort(
        key=lambda row: (
            row["layout"],
            row["seat_section"],
            row["seat_table"],
            row["seat_side"],
            int(row["seat_position"]),
            row["seat_id"],
        )
    )

    return rows


def write_solution_csvs(sets, layouts, main_seat, dessert_seat):
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    main_rows = solution_rows(
        sets,
        layouts,
        main_seat,
        "main_course",
    )

    dessert_rows = solution_rows(
        sets,
        layouts,
        dessert_seat,
        "dessert",
    )

    with MAIN_OUTPUT_CSV.open("w", encoding="utf-8", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=list(main_rows[0]))
        writer.writeheader()
        writer.writerows(main_rows)

    with DESSERT_OUTPUT_CSV.open("w", encoding="utf-8", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=list(dessert_rows[0]))
        writer.writeheader()
        writer.writerows(dessert_rows)


def load_all_data():
    diners, fellow_guest_pairs = load_diners()
    layouts = load_layouts()
    sets = derive_sets(diners, fellow_guest_pairs)
    conflicts = derive_conflicts(diners, sets)
    stage_sets = build_stage_sets(layouts)

    return diners, layouts, sets, conflicts, stage_sets


def main():
    parser = argparse.ArgumentParser(
        description="College Feast three-stage Xpress MILP."
    )
    parser.add_argument(
        "--validate-data",
        action="store_true",
        help="Only parse and validate the workbook/CSV inputs.",
    )
    parser.add_argument(
        "--stage1-only",
        action="store_true",
        help="Solve Stage 1 and write section/room assignments only.",
    )
    args = parser.parse_args()

    diners, layouts, sets, conflicts, stage_sets = load_all_data()

    validate_data(
        diners,
        layouts,
        sets,
        conflicts,
        stage_sets,
    )

    if args.validate_data:
        return

    xp = import_xpress()

    main_assignment, dessert_assignment = solve_stage_1(
        xp,
        sets,
        layouts,
        stage_sets,
    )

    write_stage1_csv(
        sets,
        main_assignment,
        dessert_assignment,
    )

    print("Stage 1 assignments written to:", STAGE1_OUTPUT_CSV)

    if args.stage1_only:
        return

    main_seat = solve_stage_2(
        xp,
        sets,
        conflicts,
        layouts,
        stage_sets,
        main_assignment,
    )

    dessert_seat, repeated_pairs = solve_stage_3(
        xp,
        sets,
        conflicts,
        layouts,
        stage_sets,
        dessert_assignment,
        main_seat,
    )

    write_solution_csvs(
        sets,
        layouts,
        main_seat,
        dessert_seat,
    )

    print("Main Course solution written to:", MAIN_OUTPUT_CSV)
    print("Dessert solution written to:", DESSERT_OUTPUT_CSV)
    print("Repeated Main Course close pairs considered:", len(repeated_pairs))


if __name__ == "__main__":
    main()