import csv
import re
from pathlib import Path
from collections import Counter, defaultdict

import openpyxl
import xpress as xp


WORKBOOK = Path(
    "/Users/yifeiwang/Desktop/dissertation/data/"
    "MSc project workbook.xlsx"
)

CSV_DIR = Path(
    "/Users/yifeiwang/Desktop/dissertation/code/csv data"
)

SEATS_CSV = CSV_DIR / "all_layouts_seats.csv"
PAIRS_CSV = CSV_DIR / "all_layouts_seat_pairs.csv"

OUTPUT_CSV = Path(
    "/Users/yifeiwang/Desktop/dissertation/code/"
    "subject_dinner_solution_single_stage.csv"
)

LAYOUT = "hall_110"
SHEET = "Subject Dinner"

LAMBDA_1_YEAR = 1.0
LAMBDA_2_GENDER = 1.0
LAMBDA_3_FELLOWS = 1.0
LAMBDA_4_CATEGORY = 1.0
LAMBDA_5_ALUMNI_STUDENT = 1.0

TIME_LIMIT_SECONDS = 25200
MIP_RELATIVE_GAP = 0.01


def clean(value):
    """
    cleans the value so blanks and extra spaces do not mess up the data.
    """
    if value is None:
        return ""

    return re.sub(r"\s+", " ", str(value)).strip()


def slug(value):
    """
    turns text into a simple lower case label for ids and variable names.
    """
    text = clean(value).lower()
    text = re.sub(r"[^a-z0-9]+", "_", text).strip("_")

    return text or "x"


def to_int_or_none(value):
    """
    tries to change the value into an integer, and returns none if it cannot.
    """
    if value is None:
        return None

    if isinstance(value, int):
        return value

    if isinstance(value, float) and value.is_integer():
        return int(value)

    text = clean(value)

    if text.isdigit():
        return int(text)

    return None


def determine_gender(title, gender):
    """
    works out the gender code by checking the title first and then the gender column.
    """
    title_value = clean(title).lower().rstrip(".")
    gender_value = clean(gender).lower()

    # Explicit titles take priority.
    if title_value in {"mr", "sir", "lord"}:
        return "M"

    if title_value in {
        "mrs",
        "miss",
        "ms",
        "madam",
        "dame",
        "lady",
    }:
        return "W"

    # Neutral titles use column C.
    if gender_value in {"m", "male"}:
        return "M"

    # Blank represents female in this workbook.
    if gender_value in {
        "",
        "f",
        "female",
        "w",
        "woman",
    }:
        return "W"

    raise ValueError(
        f"Cannot determine gender: "
        f"title={title!r}, gender={gender!r}"
    )


def load_diners():
    """
    loads the valid diners from the workbook and keeps the information needed for the model.
    """
    workbook = openpyxl.load_workbook(
        WORKBOOK,
        data_only=True,
    )

    worksheet = workbook[SHEET]

    diners = []
    section = None

    section_names = {
        "ALUMNI",
        "STUDENTS",
        "FELLOWS",
    }

    for excel_row in range(4, worksheet.max_row + 1):
        values = [
            worksheet.cell(excel_row, column).value
            for column in range(1, 9)
        ]

        (
            number,
            title,
            gender,
            first,
            surname,
            matric,
            dinner,
            notes,
        ) = values

        # Section labels may appear in column A or B.
        detected_section = None

        for value in values[:4]:
            value_text = clean(value).upper()

            if value_text in section_names:
                detected_section = value_text
                break

        if detected_section is not None:
            section = detected_section
            continue

        if section is None:
            continue

        if (
            number is None
            and first is None
            and surname is None
        ):
            continue

        if to_int_or_none(dinner) is None:
            continue

        number_integer = to_int_or_none(number)

        if number_integer is None:
            number_integer = excel_row

        gender_code = determine_gender(
            title,
            gender,
        )

        diner_id = (
            f"{section[0]}"
            f"{number_integer:03d}_"
            f"{slug(first)}_{slug(surname)}"
        )

        full_name = " ".join(
            part
            for part in [
                clean(title),
                clean(first),
                clean(surname),
            ]
            if part
        )

        diners.append({
            "id": diner_id,
            "excel_row": excel_row,
            "category": section,
            "no": number_integer,
            "title": clean(title),
            "first": clean(first),
            "surname": clean(surname),
            "name": full_name,
            "matric": to_int_or_none(matric),
            "notes": clean(notes),
            "gender": gender_code,
        })

    category_counts = Counter(
        diner["category"]
        for diner in diners
    )

    if set(category_counts) != section_names:
        raise ValueError(
            "Not all categories were detected. "
            f"Detected: {category_counts}"
        )

    return diners


def load_layout():
    """
    loads the hall seating layout and the close-seat pairs from the csv files.
    """
    seats = []
    sections = defaultdict(list)

    high_table = set()
    high_end = set()

    with SEATS_CSV.open(
        encoding="utf-8-sig",
        newline="",
    ) as file:
        reader = csv.DictReader(file)

        for row in reader:
            if row["layout"] != LAYOUT:
                continue

            seat_id = row["seat_id"]
            section_id = row["section_id"]

            seats.append({
                "id": seat_id,
                "section": section_id,
                "table": row["table_id"],
                "side": row["side"],
                "position": int(row["position"]),
                "excel_row": int(row["excel_row"]),
                "excel_col": int(row["excel_col"]),
            })

            sections[section_id].append(seat_id)

            if clean(row["is_high_table"]) == "1":
                high_table.add(seat_id)

            if clean(row["high_table_zone"]).upper() == "END":
                high_end.add(seat_id)

    if not seats:
        raise ValueError(
            "No seats were loaded for hall_110."
        )

    master_seat = "I8"

    if master_seat not in high_table:
        raise ValueError(
            f"Master seat {master_seat} "
            "is not a High Table seat."
        )

    high_centre = {
        master_seat,
    }

    high_vip = {
        "H8",
        "K8",
        "H12",
        "I12",
        "J12",
        "K12",
    }

    missing_vip = high_vip - high_table

    if missing_vip:
        raise ValueError(
            "These VIP seats are not in H: "
            f"{sorted(missing_vip)}"
        )

    if master_seat in high_vip:
        raise ValueError(
            "HVIP must exclude the Master's seat."
        )

    seat_ids = {
        seat["id"]
        for seat in seats
    }

    close_weights = {}

    with PAIRS_CSV.open(
        encoding="utf-8-sig",
        newline="",
    ) as file:
        reader = csv.DictReader(file)

        for row in reader:
            if row["layout"] != LAYOUT:
                continue

            seat_s = row["seat_s"]
            seat_r = row["seat_r"]

            if seat_s == seat_r:
                continue

            if (
                seat_s not in seat_ids
                or seat_r not in seat_ids
            ):
                continue

            # The CSV contains both directions.
            representative_pair = tuple(
                sorted((seat_s, seat_r))
            )

            weight = float(row["weight"])

            if representative_pair in close_weights:
                close_weights[representative_pair] = max(
                    close_weights[representative_pair],
                    weight,
                )
            else:
                close_weights[representative_pair] = weight

    return (
        seats,
        dict(sections),
        close_weights,
        high_table,
        high_vip,
        high_centre,
        high_end,
        master_seat,
    )


def derive_sets(diners):
    """
    makes the diner groups used in the optimisation model.
    """
    I = [
        diner["id"]
        for diner in diners
    ]

    M = {
        diner["id"]
        for diner in diners
        if diner["gender"] == "M"
    }

    W = {
        diner["id"]
        for diner in diners
        if diner["gender"] == "W"
    }

    F = {
        diner["id"]
        for diner in diners
        if diner["category"] == "FELLOWS"
    }

    A = {
        diner["id"]
        for diner in diners
        if diner["category"] == "ALUMNI"
    }

    U = {
        diner["id"]
        for diner in diners
        if diner["category"] == "STUDENTS"
    }

    IY = {
        diner["id"]
        for diner in diners
        if diner["matric"] is not None
    }

    IH = set()
    IVIP = set()
    ICENTRE = set()
    IEND = set()

    for diner in diners:
        diner_id = diner["id"]
        note = diner["notes"].lower()

        if (
            "high table" in note
            or "honor student" in note
            or "honour student" in note
        ):
            IH.add(diner_id)

        if (
            "vip" in note
            and "high" in note
        ):
            IVIP.add(diner_id)
            IH.add(diner_id)

        if (
            ("centre" in note or "center" in note)
            and "high" in note
        ):
            ICENTRE.add(diner_id)
            IH.add(diner_id)

        if (
            "end" in note
            and "high" in note
        ):
            IEND.add(diner_id)
            IH.add(diner_id)

    return (
        I,
        M,
        W,
        F,
        A,
        U,
        IY,
        IH,
        IVIP,
        ICENTRE,
        IEND,
    )


def find_master(diners, centre_diners):
    """
    finds the master diner from the name or notes, or uses the centre high table diner.
    """
    master_candidates = []

    for diner in diners:
        text = clean(
            f"{diner['title']} "
            f"{diner['first']} "
            f"{diner['surname']} "
            f"{diner['notes']}"
        ).lower()

        if "master" in text:
            master_candidates.append(
                diner["id"]
            )

    if len(master_candidates) == 1:
        return master_candidates[0]

    centre_candidates = sorted(
        centre_diners
    )

    if len(centre_candidates) == 1:
        return centre_candidates[0]

    raise ValueError(
        "Could not identify exactly one Master. "
        f"Master candidates: {master_candidates}; "
        f"centre candidates: {centre_candidates}"
    )


def main():
    """
    runs the compact seating model and writes the final seating plan to a csv file.
    """
    diners = load_diners()

    (
        seats,
        sections,
        close_weights,
        H,
        HVIP,
        HCENTRE,
        HEND,
        master_seat,
    ) = load_layout()

    (
        I,
        M,
        W,
        F,
        A,
        U,
        IY,
        IH,
        IVIP,
        ICENTRE,
        IEND,
    ) = derive_sets(diners)

    master_id = find_master(
        diners,
        ICENTRE,
    )

    S = [
        seat["id"]
        for seat in seats
    ]

    Q = list(sections)
    C = list(close_weights)

    diner_by_id = {
        diner["id"]: diner
        for diner in diners
    }

    seat_by_id = {
        seat["id"]: seat
        for seat in seats
    }

    matric = {
        diner["id"]: diner["matric"]
        for diner in diners
    }

    if not IY:
        raise ValueError(
            "No known matriculation years were found."
        )

    m_min = min(
        matric[diner_id]
        for diner_id in IY
    )

    delta_y = (
        max(
            matric[diner_id]
            for diner_id in IY
        )
        - m_min
    )

    print("Diners |I|:", len(I))
    print("Seats |S|:", len(S))
    print("Sections |Q|:", len(Q))
    print(
        "Representative close pairs |C|:",
        len(C),
    )
    print(
        "High Table seats |H|:",
        len(H),
    )
    print(
        "High Table diners |IH|:",
        len(IH),
    )
    print("Master diner:", master_id)
    print("Master seat:", master_seat)
    print("m_min:", m_min)
    print("DeltaY:", delta_y)

    print(
        "Category counts:",
        Counter(
            diner["category"]
            for diner in diners
        ),
    )

    print(
        "Gender counts:",
        Counter(
            diner["gender"]
            for diner in diners
        ),
    )

    if len(I) > len(S):
        raise ValueError(
            "There are more diners than seats."
        )

    if len(IH) > len(H):
        raise ValueError(
            "There are more High Table diners "
            "than High Table seats."
        )

    license_path = Path(
        "/Applications/FICO Xpress/"
        "xpressmp/bin/xpauth.xpr"
    )

    if license_path.exists():
        xp.init(str(license_path))

    problem = xp.problem(
        "subject_dinner_compact_tight_milp"
    )

    problem.controls.outputlog = 1
    problem.controls.timelimit = TIME_LIMIT_SECONDS
    problem.controls.miprelstop = MIP_RELATIVE_GAP

    # Assignment variables.
    x = {}

    for diner_id in I:
        for seat_id in S:
            x[diner_id, seat_id] = xp.var(
                vartype=xp.binary,
                name=f"x_{diner_id}_{seat_id}",
            )

    problem.addVariable(
        list(x.values())
    )

    # Tight Alumni-Student variables.
    z_AU_forward = {}
    z_AU_reverse = {}
    delta = {}

    for seat_s, seat_r in C:
        z_AU_forward[
            seat_s,
            seat_r,
        ] = xp.var(
            vartype=xp.binary,
            name=(
                f"z_AU_forward_"
                f"{seat_s}_{seat_r}"
            ),
        )

        z_AU_reverse[
            seat_s,
            seat_r,
        ] = xp.var(
            vartype=xp.binary,
            name=(
                f"z_AU_reverse_"
                f"{seat_s}_{seat_r}"
            ),
        )

        delta[
            seat_s,
            seat_r,
        ] = xp.var(
            lb=0,
            name=f"delta_{seat_s}_{seat_r}",
        )

    problem.addVariable(
        list(z_AU_forward.values())
        + list(z_AU_reverse.values())
        + list(delta.values())
    )

    # Section-deviation variables.
    groups = {
        "M": M,
        "W": W,
        "F": F,
        "A": A,
        "U": U,
    }

    dplus = {}
    dminus = {}

    for group_name in groups:
        for section in Q:
            dplus[
                group_name,
                section,
            ] = xp.var(
                lb=0,
                name=(
                    f"dplus_{group_name}_"
                    f"{slug(section)}"
                ),
            )

            dminus[
                group_name,
                section,
            ] = xp.var(
                lb=0,
                name=(
                    f"dminus_{group_name}_"
                    f"{slug(section)}"
                ),
            )

    problem.addVariable(
        list(dplus.values())
        + list(dminus.values())
    )

    # Every diner receives exactly one seat.
    for diner_id in I:
        problem.addConstraint(
            xp.Sum(
                x[diner_id, seat_id]
                for seat_id in S
            ) == 1
        )

    # Every seat contains at most one diner.
    for seat_id in S:
        problem.addConstraint(
            xp.Sum(
                x[diner_id, seat_id]
                for diner_id in I
            ) <= 1
        )

    # High Table requirements.
    for diner_id in IH:
        problem.addConstraint(
            xp.Sum(
                x[diner_id, seat_id]
                for seat_id in H
            ) == 1
        )

    for diner_id in IVIP:
        problem.addConstraint(
            xp.Sum(
                x[diner_id, seat_id]
                for seat_id in HVIP
            ) == 1
        )

    for diner_id in ICENTRE:
        problem.addConstraint(
            xp.Sum(
                x[diner_id, seat_id]
                for seat_id in HCENTRE
            ) == 1
        )

    for diner_id in IEND:
        problem.addConstraint(
            xp.Sum(
                x[diner_id, seat_id]
                for seat_id in HEND
            ) == 1
        )

    problem.addConstraint(
        x[master_id, master_seat] == 1
    )

    # Section balance.
    for section in Q:
        section_seats = sections[section]

        for group_name, group_set in groups.items():
            target = (
                len(group_set)
                * len(section_seats)
                / len(S)
            )

            actual = xp.Sum(
                x[diner_id, seat_id]
                for diner_id in group_set
                for seat_id in section_seats
            )

            problem.addConstraint(
                actual - target
                == dplus[group_name, section]
                - dminus[group_name, section]
            )

    # Compact seat-level expressions.
    alumni_at_seat = {}
    student_at_seat = {}
    known_year_at_seat = {}
    shifted_year_at_seat = {}

    for seat_id in S:
        alumni_at_seat[seat_id] = xp.Sum(
            x[diner_id, seat_id]
            for diner_id in A
        )

        student_at_seat[seat_id] = xp.Sum(
            x[diner_id, seat_id]
            for diner_id in U
        )

        known_year_at_seat[seat_id] = xp.Sum(
            x[diner_id, seat_id]
            for diner_id in IY
        )

        shifted_year_at_seat[seat_id] = xp.Sum(
            (
                matric[diner_id] - m_min
            )
            * x[diner_id, seat_id]
            for diner_id in IY
        )

    # Tight close-pair constraints.
    for seat_s, seat_r in C:
        # Forward: Alumnus at s and Student at r.
        problem.addConstraint(
            z_AU_forward[seat_s, seat_r]
            <= alumni_at_seat[seat_s]
        )

        problem.addConstraint(
            z_AU_forward[seat_s, seat_r]
            <= student_at_seat[seat_r]
        )

        problem.addConstraint(
            z_AU_forward[seat_s, seat_r]
            >= alumni_at_seat[seat_s]
            + student_at_seat[seat_r]
            - 1
        )

        # Reverse: Student at s and Alumnus at r.
        problem.addConstraint(
            z_AU_reverse[seat_s, seat_r]
            <= student_at_seat[seat_s]
        )

        problem.addConstraint(
            z_AU_reverse[seat_s, seat_r]
            <= alumni_at_seat[seat_r]
        )

        problem.addConstraint(
            z_AU_reverse[seat_s, seat_r]
            >= student_at_seat[seat_s]
            + alumni_at_seat[seat_r]
            - 1
        )

        problem.addConstraint(
            z_AU_forward[seat_s, seat_r]
            + z_AU_reverse[seat_s, seat_r]
            <= 1
        )

        # Matriculation-year difference.
        problem.addConstraint(
            delta[seat_s, seat_r]
            >= shifted_year_at_seat[seat_s]
            - shifted_year_at_seat[seat_r]
            - delta_y
            * (
                2
                - known_year_at_seat[seat_s]
                - known_year_at_seat[seat_r]
            )
        )

        problem.addConstraint(
            delta[seat_s, seat_r]
            >= shifted_year_at_seat[seat_r]
            - shifted_year_at_seat[seat_s]
            - delta_y
            * (
                2
                - known_year_at_seat[seat_s]
                - known_year_at_seat[seat_r]
            )
        )

        problem.addConstraint(
            delta[seat_s, seat_r]
            <= delta_y
            * known_year_at_seat[seat_s]
        )

        problem.addConstraint(
            delta[seat_s, seat_r]
            <= delta_y
            * known_year_at_seat[seat_r]
        )


    # NORMALISED OBJECTIVE
    total_pair_weight = sum(
        close_weights[seat_s, seat_r]
        for seat_s, seat_r in C
    )

    if total_pair_weight <= 0:
        raise ValueError(
            "There is no positive close-seat weight."
        )

    gender_scale = 2.0 * (len(M) + len(W))
    fellows_scale = 2.0 * len(F)
    category_scale = 2.0 * (len(A) + len(U))

    normalised_year_term = (
        xp.Sum(
            close_weights[seat_s, seat_r]
            * delta[seat_s, seat_r]
            for seat_s, seat_r in C
        )
        / (
            delta_y
            * total_pair_weight
        )
    )

    normalised_gender_term = (
        xp.Sum(
            dplus["M", section]
            + dminus["M", section]
            + dplus["W", section]
            + dminus["W", section]
            for section in Q
        )
        / gender_scale
    )

    normalised_fellows_term = (
        xp.Sum(
            dplus["F", section]
            + dminus["F", section]
            for section in Q
        )
        / fellows_scale
    )

    normalised_category_term = (
        xp.Sum(
            dplus["A", section]
            + dminus["A", section]
            + dplus["U", section]
            + dminus["U", section]
            for section in Q
        )
        / category_scale
    )

    normalised_alumni_student_term = (
        xp.Sum(
            close_weights[seat_s, seat_r]
            * (
                z_AU_forward[seat_s, seat_r]
                + z_AU_reverse[seat_s, seat_r]
            )
            for seat_s, seat_r in C
        )
        / total_pair_weight
    )

    objective = (
        LAMBDA_1_YEAR
        * normalised_year_term
        + LAMBDA_2_GENDER
        * normalised_gender_term
        + LAMBDA_3_FELLOWS
        * normalised_fellows_term
        + LAMBDA_4_CATEGORY
        * normalised_category_term
        - LAMBDA_5_ALUMNI_STUDENT
        * normalised_alumni_student_term
    )

    problem.setObjective(
        objective,
        sense=xp.minimize,
    )

    print("Tight compact model built. Solving...")

    problem.solve()

    print(
        "Solve status:",
        problem.attributes.solvestatus,
    )
    print(
        "Solution status:",
        problem.attributes.solstatus,
    )
    print(
        "Number of MIP solutions:",
        problem.attributes.mipsols,
    )

    if problem.attributes.mipsols < 1:
        raise RuntimeError(
            "No integer solution is available."
        )

    print(
        "Objective value:",
        problem.attributes.objval,
    )

    output_rows = []

    for diner_id in I:
        assigned_seat = None

        for seat_id in S:
            if (
                problem.getSolution(
                    x[diner_id, seat_id]
                )
                > 0.5
            ):
                assigned_seat = seat_id
                break

        if assigned_seat is None:
            raise RuntimeError(
                f"Diner {diner_id} has no assigned seat."
            )

        diner = diner_by_id[diner_id]
        seat = seat_by_id[assigned_seat]

        output_rows.append({
            "seat_id": assigned_seat,
            "seat_section": seat["section"],
            "seat_table": seat["table"],
            "seat_side": seat["side"],
            "seat_position": seat["position"],
            "diner_id": diner_id,
            "category": diner["category"],
            "title": diner["title"],
            "first": diner["first"],
            "surname": diner["surname"],
            "full_name": diner["name"],
            "matric": (
                diner["matric"]
                if diner["matric"] is not None
                else ""
            ),
            "gender": diner["gender"],
            "notes": diner["notes"],
        })

    if (
        len({
            row["seat_id"]
            for row in output_rows
        })
        != len(output_rows)
    ):
        raise RuntimeError(
            "More than one diner was assigned "
            "to the same seat."
        )

    output_rows.sort(
        key=lambda row: (
            row["seat_section"],
            row["seat_table"],
            row["seat_side"],
            int(row["seat_position"]),
        )
    )

    fieldnames = [
        "seat_id",
        "seat_section",
        "seat_table",
        "seat_side",
        "seat_position",
        "diner_id",
        "category",
        "title",
        "first",
        "surname",
        "full_name",
        "matric",
        "gender",
        "notes",
    ]

    with OUTPUT_CSV.open(
        "w",
        encoding="utf-8",
        newline="",
    ) as file:
        writer = csv.DictWriter(
            file,
            fieldnames=fieldnames,
        )

        writer.writeheader()
        writer.writerows(
            output_rows
        )

    print(
        "Solution written to:",
        OUTPUT_CSV,
    )


if __name__ == "__main__":
    main()