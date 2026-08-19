import csv
import re
from pathlib import Path
from collections import Counter, defaultdict

import openpyxl
import xpress as xp



# FILE PATHS
WORKBOOK = Path(
    "/Users/yifeiwang/Desktop/dissertation/data/"
    "MSc project workbook.xlsx"
)

CSV_DIR = Path(
    "/Users/yifeiwang/Desktop/dissertation/code/csv data"
)

SEATS_CSV = CSV_DIR / "all_layouts_seats.csv"
PAIRS_CSV = CSV_DIR / "all_layouts_seat_pairs.csv"

OUTPUT_CSV = Path(
    "/Users/yifeiwang/Desktop/dissertation/code/"
    "subject_dinner_solution_two_stage.csv"
)

LAYOUT = "hall_110"
SHEET = "Subject Dinner"



# OBJECTIVE WEIGHTS AND SOLVER SETTINGS
ALPHA_1_GENDER = 1.0
ALPHA_2_FELLOWS = 1.0
ALPHA_3_CATEGORY = 1.0
ALPHA_4_YEAR_GROUP_SPREAD = 1.0

NUM_YEAR_GROUPS = 5

BETA_1_YEAR = 1.0
BETA_2_ALUMNI_STUDENT = 1.0

STAGE_1_TIME_LIMIT = 600
STAGE_2_TIME_LIMIT = 3600

MIP_RELATIVE_GAP = 0.01



# STAGE 1 SWITCHES
USE_HIGH_TABLE = True
USE_GENDER_BALANCE = True
USE_FELLOWS_BALANCE = True
USE_CATEGORY_BALANCE = True
USE_YEAR_GROUP_SPREAD = True

RUN_STAGE_2 = True



# GENERAL FUNCTIONS
def clean(value):
    """
    cleans the value from excel so it is easier to use later.
    """
    if value is None:
        return ""

    return re.sub(r"\s+", " ", str(value)).strip()


def slug(value):
    """
    turns a name or text into a simple lower case code for ids and variables.
    """
    text = clean(value).lower()
    text = re.sub(r"[^a-z0-9]+", "_", text).strip("_")

    return text or "x"


def to_int_or_none(value):
    """
    tries to change the value into an int, and gives none if it cannot.
    """
    if value is None:
        return None

    if isinstance(value, int):
        return value

    if isinstance(value, float) and value.is_integer():
        return int(value)

    text = clean(value)

    if text.isdigit():
        return int(text)

    return None

def build_year_groups(
    known_year_diners,
    matric,
    num_groups,
):
    """
    splits diners into similar sized matriculation year groups and keeps
    people from the same year together.
    """
    if num_groups < 2:
        raise ValueError(
            "NUM_YEAR_GROUPS must be at least 2."
        )

    diners_by_year = defaultdict(list)

    for diner_id in known_year_diners:
        year = matric[diner_id]
        diners_by_year[year].append(diner_id)

    unique_years = sorted(diners_by_year)

    if len(unique_years) < num_groups:
        raise ValueError(
            "There are fewer distinct matriculation years "
            f"({len(unique_years)}) than requested groups "
            f"({num_groups})."
        )

    total_diners = len(known_year_diners)

    year_groups = {
        group_id: set()
        for group_id in range(1, num_groups + 1)
    }

    cutoffs = []
    current_group = 1
    cumulative_count = 0

    for position, year in enumerate(unique_years):
        year_diners = diners_by_year[year]

        year_groups[current_group].update(
            year_diners
        )

        cumulative_count += len(year_diners)

        remaining_years = (
            len(unique_years) - position - 1
        )

        remaining_groups = (
            num_groups - current_group
        )

        target_cumulative_count = (
            current_group
            * total_diners
            / num_groups
        )

        reached_target = (
            cumulative_count
            >= target_cumulative_count
        )

        must_start_next_group = (
            remaining_years
            == remaining_groups
        )

        if (
            current_group < num_groups
            and (
                reached_target
                or must_start_next_group
            )
        ):
            cutoffs.append(year)
            current_group += 1

    for group_id, group_diners in year_groups.items():
        if not group_diners:
            raise ValueError(
                f"Year group {group_id} is empty."
            )

    if len(cutoffs) != num_groups - 1:
        raise ValueError(
            "Could not create the requested number "
            "of matriculation-year groups."
        )

    return cutoffs, year_groups


def determine_gender(title, gender):
    """
    works out the gender code by checking the title first, then the gender column.
    """
    title_value = clean(title).lower().rstrip(".")
    gender_value = clean(gender).lower()

    if title_value in {"mr", "sir", "lord"}:
        return "M"

    if title_value in {
        "mrs",
        "miss",
        "ms",
        "madam",
        "dame",
        "lady",
    }:
        return "W"

    if gender_value in {"m", "male"}:
        return "M"

    if gender_value in {
        "",
        "f",
        "female",
        "w",
        "woman",
    }:
        return "W"

    raise ValueError(
        f"Cannot determine gender: "
        f"title={title!r}, gender={gender!r}"
    )



# LOAD DINERS
def load_diners():
    workbook = openpyxl.load_workbook(
        WORKBOOK,
        data_only=True,
    )
    """
    loads all valid diners from the workbook and stores the useful details for
    the seating model.
    """
    worksheet = workbook[SHEET]

    diners = []
    current_section = None

    section_names = {
        "ALUMNI",
        "STUDENTS",
        "FELLOWS",
    }

    for excel_row in range(4, worksheet.max_row + 1):
        values = [
            worksheet.cell(excel_row, column).value
            for column in range(1, 9)
        ]

        (
            number,
            title,
            gender,
            first,
            surname,
            matric,
            dinner,
            notes,
        ) = values

        detected_section = None

        for value in values[:4]:
            value_text = clean(value).upper()

            if value_text in section_names:
                detected_section = value_text
                break

        if detected_section is not None:
            current_section = detected_section
            continue

        if current_section is None:
            continue

        if (
            number is None
            and first is None
            and surname is None
        ):
            continue

        if to_int_or_none(dinner) is None:
            continue

        number_integer = to_int_or_none(number)

        if number_integer is None:
            number_integer = excel_row

        gender_code = determine_gender(
            title,
            gender,
        )

        diner_id = (
            f"{current_section[0]}"
            f"{number_integer:03d}_"
            f"{slug(first)}_{slug(surname)}"
        )

        full_name = " ".join(
            part
            for part in [
                clean(title),
                clean(first),
                clean(surname),
            ]
            if part
        )

        diners.append({
            "id": diner_id,
            "excel_row": excel_row,
            "category": current_section,
            "no": number_integer,
            "title": clean(title),
            "first": clean(first),
            "surname": clean(surname),
            "name": full_name,
            "matric": to_int_or_none(matric),
            "notes": clean(notes),
            "gender": gender_code,
        })

    category_counts = Counter(
        diner["category"]
        for diner in diners
    )

    if set(category_counts) != section_names:
        raise ValueError(
            "Not all diner categories were detected. "
            f"Detected counts: {category_counts}"
        )

    return diners



# LOAD SEAT LAYOUT
def load_layout():
    """
    loads the seat layout from the csv files
    and also finds the special high table seats.
    """
    seats = []
    sections = defaultdict(list)

    high_table = set()
    high_end = set()

    with SEATS_CSV.open(
        encoding="utf-8-sig",
        newline="",
    ) as file:
        reader = csv.DictReader(file)

        for row in reader:
            if row["layout"] != LAYOUT:
                continue

            seat_id = row["seat_id"]
            section_id = row["section_id"]

            seats.append({
                "id": seat_id,
                "section": section_id,
                "table": row["table_id"],
                "side": row["side"],
                "position": int(row["position"]),
                "excel_row": int(row["excel_row"]),
                "excel_col": int(row["excel_col"]),
            })

            sections[section_id].append(seat_id)

            if clean(row["is_high_table"]) == "1":
                high_table.add(seat_id)

            if clean(row["high_table_zone"]).upper() == "END":
                high_end.add(seat_id)

    if not seats:
        raise ValueError(
            "No seats were loaded for hall_110."
        )

    master_seat = "I8"

    if master_seat not in high_table:
        raise ValueError(
            f"Master seat {master_seat} "
            "is not a High Table seat."
        )

    high_centre = {
        master_seat,
    }

    high_vip = {
        "H8",
        "K8",
        "H12",
        "I12",
        "J12",
        "K12",
    }

    missing_vip = high_vip - high_table

    if missing_vip:
        raise ValueError(
            "These VIP seats are not High Table seats: "
            f"{sorted(missing_vip)}"
        )

    if master_seat in high_vip:
        raise ValueError(
            "The Master's seat must not be in HVIP."
        )

    seat_ids = {
        seat["id"]
        for seat in seats
    }

    close_weights = {}

    with PAIRS_CSV.open(
        encoding="utf-8-sig",
        newline="",
    ) as file:
        reader = csv.DictReader(file)

        for row in reader:
            if row["layout"] != LAYOUT:
                continue

            seat_s = row["seat_s"]
            seat_r = row["seat_r"]

            if seat_s == seat_r:
                continue

            if (
                seat_s not in seat_ids
                or seat_r not in seat_ids
            ):
                continue

            representative_pair = tuple(
                sorted((seat_s, seat_r))
            )

            weight = float(row["weight"])

            if representative_pair in close_weights:
                close_weights[representative_pair] = max(
                    close_weights[representative_pair],
                    weight,
                )
            else:
                close_weights[representative_pair] = weight

    return (
        seats,
        dict(sections),
        close_weights,
        high_table,
        high_vip,
        high_centre,
        high_end,
        master_seat,
    )



# DERIVE SETS
def derive_sets(diners):
    """
    makes the main diner groups that are needed in the optimisation model.
    """
    I = [
        diner["id"]
        for diner in diners
    ]

    M = {
        diner["id"]
        for diner in diners
        if diner["gender"] == "M"
    }

    W = {
        diner["id"]
        for diner in diners
        if diner["gender"] == "W"
    }

    F = {
        diner["id"]
        for diner in diners
        if diner["category"] == "FELLOWS"
    }

    A = {
        diner["id"]
        for diner in diners
        if diner["category"] == "ALUMNI"
    }

    U = {
        diner["id"]
        for diner in diners
        if diner["category"] == "STUDENTS"
    }

    IY = {
        diner["id"]
        for diner in diners
        if diner["matric"] is not None
    }

    IH = set()
    IVIP = set()
    ICENTRE = set()
    IEND = set()

    for diner in diners:
        diner_id = diner["id"]
        note = diner["notes"].lower()

        if (
            "high table" in note
            or "honor student" in note
            or "honour student" in note
        ):
            IH.add(diner_id)

        if (
            "vip" in note
            and "high" in note
        ):
            IVIP.add(diner_id)
            IH.add(diner_id)

        if (
            ("centre" in note or "center" in note)
            and "high" in note
        ):
            ICENTRE.add(diner_id)
            IH.add(diner_id)

        if (
            "end" in note
            and "high" in note
        ):
            IEND.add(diner_id)
            IH.add(diner_id)

    return (
        I,
        M,
        W,
        F,
        A,
        U,
        IY,
        IH,
        IVIP,
        ICENTRE,
        IEND,
    )


def find_master(diners, centre_diners):
    """
    finds the master diner, either from the text fields or from the centre
    high table group.
    """
    master_candidates = []

    for diner in diners:
        text = clean(
            f"{diner['title']} "
            f"{diner['first']} "
            f"{diner['surname']} "
            f"{diner['notes']}"
        ).lower()

        if "master" in text:
            master_candidates.append(
                diner["id"]
            )

    if len(master_candidates) == 1:
        return master_candidates[0]

    centre_candidates = sorted(
        centre_diners
    )

    if len(centre_candidates) == 1:
        return centre_candidates[0]

    raise ValueError(
        "Could not identify exactly one Master. "
        f"Master candidates: {master_candidates}; "
        f"centre candidates: {centre_candidates}"
    )



# MAIN
def main():
    """
    runs the whole seating model, first assigning diners to sections and then
    assigning exact seats.
    """
    diners = load_diners()

    (
        seats,
        sections,
        close_weights,
        H,
        HVIP,
        HCENTRE,
        HEND,
        master_seat,
    ) = load_layout()

    (
        I,
        M,
        W,
        F,
        A,
        U,
        IY,
        IH,
        IVIP,
        ICENTRE,
        IEND,
    ) = derive_sets(diners)

    master_id = find_master(
        diners,
        ICENTRE,
    )

    S = [
        seat["id"]
        for seat in seats
    ]

    Q = list(sections)

    diner_by_id = {
        diner["id"]: diner
        for diner in diners
    }

    seat_by_id = {
        seat["id"]: seat
        for seat in seats
    }

    matric = {
        diner["id"]: diner["matric"]
        for diner in diners
    }

    if not IY:
        raise ValueError(
            "No known matriculation years were found."
        )

    m_min = min(
        matric[diner_id]
        for diner_id in IY
    )

    delta_y = (
        max(
            matric[diner_id]
            for diner_id in IY
        )
        - m_min
    )

    if delta_y <= 0:
        raise ValueError(
            "DeltaY must be positive."
        )

    year_cutoffs, year_groups = build_year_groups(
        IY,
        matric,
        NUM_YEAR_GROUPS,
    )

    q_high = seat_by_id[
        master_seat
    ]["section"]

    stage2_block_definitions = [
        ("HIGH_TABLE", ["HIGH_TABLE"]),
        ("LEFT_UPPER", ["LEFT_UPPER"]),
        (
            "LEFT_LOWER",
            [
                "LEFT_LOWER",
                "LEFT_LOWER_TOP",
                "LEFT_LOWER_BOTTOM",
            ],
        ),
        (
            "MIDDLE",
            [
                "MIDDLE_UPPER",
                "MIDDLE_LOWER",
            ],
        ),
        ("RIGHT_UPPER", ["RIGHT_UPPER"]),
        ("RIGHT_LOWER", ["RIGHT_LOWER"]),
    ]

    stage2_blocks = []
    used_sections = set()

    for block_name, candidate_sections in stage2_block_definitions:
        block_sections = [
            section
            for section in candidate_sections
            if section in sections
        ]

        if block_sections:
            stage2_blocks.append(
                (block_name, block_sections)
            )
            used_sections.update(block_sections)

    for section in Q:
        if section not in used_sections:
            stage2_blocks.append(
                (section, [section])
            )

    seat_to_stage2_block = {}

    for block_name, block_sections in stage2_blocks:
        for original_section in block_sections:
            for seat_id in sections[original_section]:
                seat_to_stage2_block[seat_id] = block_name

    C_by_block = {
        block_name: []
        for block_name, _ in stage2_blocks
    }


    for seat_s, seat_r in close_weights:
        block_s = seat_to_stage2_block[seat_s]
        block_r = seat_to_stage2_block[seat_r]

        if block_s == block_r:
            C_by_block[block_s].append(
                (seat_s, seat_r)
            )

    print("\nStage 2 blocks:")
    for block_name, block_sections in stage2_blocks:
        print(
            block_name,
            block_sections,
            "seats:",
            sum(len(sections[s]) for s in block_sections),
            "close pairs:",
            len(C_by_block[block_name]),
        )

    category_counts = Counter(
        diner["category"]
        for diner in diners
    )

    gender_counts = Counter(
        diner["gender"]
        for diner in diners
    )

    print("Diners |I|:", len(I))
    print("Seats |S|:", len(S))
    print("Sections |Q|:", len(Q))
    print(
        "Representative close pairs |C|:",
        len(close_weights),
    )
    print(
        "High Table section:",
        q_high,
    )
    print(
        "High Table diners:",
        len(IH),
    )
    print(
        "Master diner:",
        master_id,
    )
    print(
        "Master seat:",
        master_seat,
    )
    print("m_min:", m_min)
    print("DeltaY:", delta_y)
    print(
        "Number of year groups:",
        NUM_YEAR_GROUPS,
    )
    print(
        "Year group cut-offs:",
        year_cutoffs,
    )
    print(
        "Year group counts:",
        {
            group_id: len(group_diners)
            for group_id, group_diners
            in year_groups.items()
        },
    )
    print("Year group ranges:")

    for group_id, group_diners in year_groups.items():
        group_years = [
            matric[diner_id]
            for diner_id in group_diners
        ]

        print(
            group_id,
            min(group_years),
            "-",
            max(group_years),
            "count:",
            len(group_diners),
        )
    print(
        "Category counts:",
        category_counts,
    )
    print(
        "Gender counts:",
        gender_counts,
    )

    print("\nStage 1 switches:")
    print("USE_HIGH_TABLE:", USE_HIGH_TABLE)
    print("USE_GENDER_BALANCE:", USE_GENDER_BALANCE)
    print("USE_FELLOWS_BALANCE:", USE_FELLOWS_BALANCE)
    print("USE_CATEGORY_BALANCE:", USE_CATEGORY_BALANCE)
    print("USE_YEAR_GROUP_SPREAD:", USE_YEAR_GROUP_SPREAD)
    print("RUN_STAGE_2:", RUN_STAGE_2)

    if len(I) > len(S):
        raise ValueError(
            "There are more diners than seats."
        )

    if len(IH) > len(H):
        raise ValueError(
            "There are more High Table diners "
            "than High Table seats."
        )

    xpress_license = Path(
        "/Applications/FICO Xpress/"
        "xpressmp/bin/xpauth.xpr"
    )

    if xpress_license.exists():
        xp.init(
            str(xpress_license)
        )


    # STAGE 1
    stage1 = xp.problem(
        "subject_dinner_stage_1"
    )

    stage1.controls.outputlog = 1
    stage1.controls.timelimit = STAGE_1_TIME_LIMIT
    stage1.controls.miprelstop = MIP_RELATIVE_GAP

    section_choice = {}

    for diner_id in I:
        for section in Q:
            section_choice[
                diner_id,
                section,
            ] = xp.var(
                vartype=xp.binary,
                name=(
                    f"a_{diner_id}_"
                    f"{slug(section)}"
                ),
            )

    stage1.addVariable(
        list(section_choice.values())
    )

    year_group_used = {}

    if USE_YEAR_GROUP_SPREAD:
        for group_id in year_groups:
            for section in Q:
                year_group_used[
                    group_id,
                    section,
                ] = xp.var(
                    vartype=xp.binary,
                    name=(
                        f"y_year_group_"
                        f"{group_id}_"
                        f"{slug(section)}"
                    ),
                )

        stage1.addVariable(
            list(year_group_used.values())
        )

    active_groups = {}

    if USE_GENDER_BALANCE:
        active_groups["M"] = M
        active_groups["W"] = W

    if USE_FELLOWS_BALANCE:
        active_groups["F"] = F

    if USE_CATEGORY_BALANCE:
        active_groups["A"] = A
        active_groups["U"] = U

    dplus = {}
    dminus = {}

    for group_name in active_groups:
        for section in Q:
            dplus[
                group_name,
                section,
            ] = xp.var(
                lb=0,
                name=(
                    f"stage1_dplus_"
                    f"{group_name}_"
                    f"{slug(section)}"
                ),
            )

            dminus[
                group_name,
                section,
            ] = xp.var(
                lb=0,
                name=(
                    f"stage1_dminus_"
                    f"{group_name}_"
                    f"{slug(section)}"
                ),
            )

    balance_variables = (
        list(dplus.values())
        + list(dminus.values())
    )

    if balance_variables:
        stage1.addVariable(
            balance_variables
        )

    # Each diner belongs to exactly one section.
    for diner_id in I:
        stage1.addConstraint(
            xp.Sum(
                section_choice[
                    diner_id,
                    section,
                ]
                for section in Q
            ) == 1
        )

    # Section capacities.
    for section in Q:
        stage1.addConstraint(
            xp.Sum(
                section_choice[
                    diner_id,
                    section,
                ]
                for diner_id in I
            ) <= len(sections[section])
        )

    # High Table section.
    if USE_HIGH_TABLE:
        for diner_id in IH:
            stage1.addConstraint(
                section_choice[
                    diner_id,
                    q_high,
                ] == 1
            )

    # Broad matriculation-year group spread constraints.
    if USE_YEAR_GROUP_SPREAD:
        for group_id, group_diners in year_groups.items():
            for section in Q:
                group_used = year_group_used[
                    group_id,
                    section,
                ]

                for diner_id in group_diners:
                    stage1.addConstraint(
                        section_choice[
                            diner_id,
                            section,
                        ] <= group_used
                    )

                stage1.addConstraint(
                    group_used
                    <= xp.Sum(
                        section_choice[
                            diner_id,
                            section,
                        ]
                        for diner_id in group_diners
                    )
                )

    # Balance constraints.
    for group_name, group_set in active_groups.items():
        for section in Q:
            target = (
                len(group_set)
                * len(sections[section])
                / len(S)
            )

            actual = xp.Sum(
                section_choice[
                    diner_id,
                    section,
                ]
                for diner_id in group_set
            )

            stage1.addConstraint(
                actual - target
                == dplus[
                    group_name,
                    section,
                ]
                - dminus[
                    group_name,
                    section,
                ]
            )


    # NORMALISED STAGE 1 OBJECTIVE
    stage1_objective = 0

    # Upper bounds used to put the four objective components
    # on approximately the same 0--1 scale.
    gender_scale = 2.0 * (len(M) + len(W))
    fellows_scale = 2.0 * len(F)
    category_scale = 2.0 * (len(A) + len(U))

    year_group_scale = (
        len(year_groups)
        * (len(Q) - 1)
    )

    if USE_GENDER_BALANCE:
        for section in Q:
            stage1_objective += (
                ALPHA_1_GENDER
                * (
                    dplus["M", section]
                    + dminus["M", section]
                    + dplus["W", section]
                    + dminus["W", section]
                )
                / gender_scale
            )

    if USE_FELLOWS_BALANCE:
        for section in Q:
            stage1_objective += (
                ALPHA_2_FELLOWS
                * (
                    dplus["F", section]
                    + dminus["F", section]
                )
                / fellows_scale
            )

    if USE_CATEGORY_BALANCE:
        for section in Q:
            stage1_objective += (
                ALPHA_3_CATEGORY
                * (
                    dplus["A", section]
                    + dminus["A", section]
                    + dplus["U", section]
                    + dminus["U", section]
                )
                / category_scale
            )

    if USE_YEAR_GROUP_SPREAD:
        stage1_objective += (
            ALPHA_4_YEAR_GROUP_SPREAD
            * (
                xp.Sum(
                    year_group_used[
                        group_id,
                        section,
                    ]
                    for group_id in year_groups
                    for section in Q
                )
                - len(year_groups)
            )
            / year_group_scale
        )

    stage1.setObjective(
        stage1_objective,
        sense=xp.minimize,
    )

    print(
        "\nStage 1 model built. Solving..."
    )

    stage1.solve()

    print(
        "Stage 1 solve status:",
        stage1.attributes.solvestatus,
    )
    print(
        "Stage 1 solution status:",
        stage1.attributes.solstatus,
    )
    print(
        "Stage 1 number of solutions:",
        stage1.attributes.mipsols,
    )

    if stage1.attributes.mipsols < 1:
        raise RuntimeError(
            "Stage 1 did not find "
            "an integer solution."
        )

    print(
        "Stage 1 objective:",
        stage1.attributes.objval,
    )

    section_assignment = {}

    for diner_id in I:
        selected_sections = [
            section
            for section in Q
            if stage1.getSolution(
                section_choice[
                    diner_id,
                    section,
                ]
            ) > 0.5
        ]

        if len(selected_sections) != 1:
            raise RuntimeError(
                f"Diner {diner_id} does not "
                "have exactly one section."
            )

        section_assignment[diner_id] = (
            selected_sections[0]
        )

    print(
        "\nStage 1 section sizes:"
    )

    for section in Q:
        section_size = sum(
            section_assignment[diner_id]
            == section
            for diner_id in I
        )

        print(
            section,
            section_size,
            "/",
            len(sections[section]),
        )

    if not RUN_STAGE_2:
        print(
            "\nRUN_STAGE_2 is False, so the program stops after Stage 1."
        )
        return


    # STAGE 2
    final_seat = {}
    stage2_objectives = {}

    for section, block_sections in stage2_blocks:
        section_diners = [
            diner_id
            for diner_id in I
            if (
                section_assignment[diner_id]
                in block_sections
            )
        ]

        section_seats = [
            seat_id
            for original_section in block_sections
            for seat_id in sections[original_section]
        ]

        section_pairs = C_by_block[section]

        section_alumni = (
            A.intersection(section_diners)
        )

        section_students = (
            U.intersection(section_diners)
        )

        section_known_year = (
            IY.intersection(section_diners)
        )

        print(
            f"\nStage 2 section {section}: "
            f"{len(section_diners)} diners, "
            f"{len(section_seats)} seats, "
            f"{len(section_pairs)} close pairs"
        )

        stage2 = xp.problem(
            f"subject_dinner_stage_2_"
            f"{slug(section)}"
        )

        stage2.controls.outputlog = 1
        stage2.controls.timelimit = (
            STAGE_2_TIME_LIMIT
        )
        stage2.controls.miprelstop = (
            MIP_RELATIVE_GAP
        )

        seat_choice = {}

        for diner_id in section_diners:
            for seat_id in section_seats:
                seat_choice[
                    diner_id,
                    seat_id,
                ] = xp.var(
                    vartype=xp.binary,
                    name=(
                        f"x_{slug(section)}_"
                        f"{diner_id}_{seat_id}"
                    ),
                )

        stage2.addVariable(
            list(seat_choice.values())
        )

        alumni_student_pair = {}
        year_difference = {}

        for seat_s, seat_r in section_pairs:
            alumni_student_pair[
                seat_s,
                seat_r,
            ] = xp.var(
                vartype=xp.binary,
                name=(
                    f"z_AU_{slug(section)}_"
                    f"{seat_s}_{seat_r}"
                ),
            )

            year_difference[
                seat_s,
                seat_r,
            ] = xp.var(
                lb=0,
                name=(
                    f"delta_{slug(section)}_"
                    f"{seat_s}_{seat_r}"
                ),
            )

        stage2.addVariable(
            list(
                alumni_student_pair.values()
            )
            + list(
                year_difference.values()
            )
        )

        # Each diner receives one seat.
        for diner_id in section_diners:
            stage2.addConstraint(
                xp.Sum(
                    seat_choice[
                        diner_id,
                        seat_id,
                    ]
                    for seat_id in section_seats
                ) == 1
            )

        # Each seat contains at most one diner.
        for seat_id in section_seats:
            stage2.addConstraint(
                xp.Sum(
                    seat_choice[
                        diner_id,
                        seat_id,
                    ]
                    for diner_id in section_diners
                ) <= 1
            )

        # Keep the Stage 1 table-part assignment.
        # Several table parts may be solved together in Stage 2
        # only so that junction close pairs such as D43--D45
        # are included in the objective.
        if len(block_sections) > 1:
            for diner_id in section_diners:
                assigned_section = section_assignment[diner_id]

                for original_section in block_sections:
                    required = (
                        1
                        if assigned_section == original_section
                        else 0
                    )

                    stage2.addConstraint(
                        xp.Sum(
                            seat_choice[
                                diner_id,
                                seat_id,
                            ]
                            for seat_id in sections[original_section]
                        ) == required
                    )

        # Exact High Table constraints.
        if q_high in block_sections:
            missing_high_diners = (
                IH - set(section_diners)
            )

            if missing_high_diners:
                raise RuntimeError(
                    "These High Table diners were "
                    "not allocated to HIGH_TABLE: "
                    f"{sorted(missing_high_diners)}"
                )

            for diner_id in IH:
                stage2.addConstraint(
                    xp.Sum(
                        seat_choice[
                            diner_id,
                            seat_id,
                        ]
                        for seat_id in H
                    ) == 1
                )

            for diner_id in IVIP:
                stage2.addConstraint(
                    xp.Sum(
                        seat_choice[
                            diner_id,
                            seat_id,
                        ]
                        for seat_id in HVIP
                    ) == 1
                )

            for diner_id in ICENTRE:
                stage2.addConstraint(
                    xp.Sum(
                        seat_choice[
                            diner_id,
                            seat_id,
                        ]
                        for seat_id in HCENTRE
                    ) == 1
                )

            for diner_id in IEND:
                stage2.addConstraint(
                    xp.Sum(
                        seat_choice[
                            diner_id,
                            seat_id,
                        ]
                        for seat_id in HEND
                    ) == 1
                )

            stage2.addConstraint(
                seat_choice[
                    master_id,
                    master_seat,
                ] == 1
            )

        alumni_at_seat = {}
        student_at_seat = {}
        known_year_at_seat = {}
        shifted_year_at_seat = {}

        for seat_id in section_seats:
            alumni_at_seat[
                seat_id
            ] = xp.Sum(
                seat_choice[
                    diner_id,
                    seat_id,
                ]
                for diner_id in section_alumni
            )

            student_at_seat[
                seat_id
            ] = xp.Sum(
                seat_choice[
                    diner_id,
                    seat_id,
                ]
                for diner_id in section_students
            )

            known_year_at_seat[
                seat_id
            ] = xp.Sum(
                seat_choice[
                    diner_id,
                    seat_id,
                ]
                for diner_id in section_known_year
            )

            shifted_year_at_seat[
                seat_id
            ] = xp.Sum(
                (
                    matric[diner_id]
                    - m_min
                )
                * seat_choice[
                    diner_id,
                    seat_id,
                ]
                for diner_id
                in section_known_year
            )

        for seat_s, seat_r in section_pairs:
            z_variable = alumni_student_pair[
                seat_s,
                seat_r,
            ]

            delta_variable = year_difference[
                seat_s,
                seat_r,
            ]

            stage2.addConstraint(
                z_variable
                <= alumni_at_seat[seat_s]
                + alumni_at_seat[seat_r]
            )

            stage2.addConstraint(
                z_variable
                <= student_at_seat[seat_s]
                + student_at_seat[seat_r]
            )

            stage2.addConstraint(
                z_variable
                >= alumni_at_seat[seat_s]
                + student_at_seat[seat_r]
                - 1
            )

            stage2.addConstraint(
                z_variable
                >= student_at_seat[seat_s]
                + alumni_at_seat[seat_r]
                - 1
            )

            stage2.addConstraint(
                delta_variable
                >= shifted_year_at_seat[seat_s]
                - shifted_year_at_seat[seat_r]
                - delta_y
                * (
                    2
                    - known_year_at_seat[seat_s]
                    - known_year_at_seat[seat_r]
                )
            )

            stage2.addConstraint(
                delta_variable
                >= shifted_year_at_seat[seat_r]
                - shifted_year_at_seat[seat_s]
                - delta_y
                * (
                    2
                    - known_year_at_seat[seat_s]
                    - known_year_at_seat[seat_r]
                )
            )

            stage2.addConstraint(
                delta_variable
                <= delta_y
                * known_year_at_seat[seat_s]
            )

            stage2.addConstraint(
                delta_variable
                <= delta_y
                * known_year_at_seat[seat_r]
            )


        # NORMALISED STAGE 2 OBJECTIVE
        total_pair_weight = sum(
            close_weights[
                seat_s,
                seat_r,
            ]
            for seat_s, seat_r
            in section_pairs
        )

        if total_pair_weight <= 0:
            raise ValueError(
                f"Section {section} has no "
                "positive close-seat weight."
            )

        normalised_year_term = (
            xp.Sum(
                close_weights[
                    seat_s,
                    seat_r,
                ]
                * year_difference[
                    seat_s,
                    seat_r,
                ]
                for seat_s, seat_r
                in section_pairs
            )
            / (
                delta_y
                * total_pair_weight
            )
        )

        normalised_alumni_student_term = (
            xp.Sum(
                close_weights[
                    seat_s,
                    seat_r,
                ]
                * alumni_student_pair[
                    seat_s,
                    seat_r,
                ]
                for seat_s, seat_r
                in section_pairs
            )
            / total_pair_weight
        )



        stage2_objective = (
            BETA_1_YEAR
            * normalised_year_term
            - BETA_2_ALUMNI_STUDENT
            * normalised_alumni_student_term
        )


        stage2.setObjective(
            stage2_objective,
            sense=xp.minimize,
        )

        print(
            f"Stage 2 model for "
            f"{section} built. Solving..."
        )

        stage2.solve()

        print(
            "Solve status:",
            stage2.attributes.solvestatus,
        )
        print(
            "Solution status:",
            stage2.attributes.solstatus,
        )
        print(
            "Number of solutions:",
            stage2.attributes.mipsols,
        )

        if stage2.attributes.mipsols < 1:
            raise RuntimeError(
                "Stage 2 did not find "
                "an integer solution for "
                f"section {section}."
            )

        stage2_objectives[section] = (
            stage2.attributes.objval
        )

        print(
            "Stage 2 objective:",
            stage2_objectives[section],
        )

        for diner_id in section_diners:
            selected_seats = [
                seat_id
                for seat_id in section_seats
                if stage2.getSolution(
                    seat_choice[
                        diner_id,
                        seat_id,
                    ]
                ) > 0.5
            ]

            if len(selected_seats) != 1:
                raise RuntimeError(
                    f"Diner {diner_id} does not "
                    f"have exactly one seat in "
                    f"{section}."
                )

            final_seat[diner_id] = (
                selected_seats[0]
            )

    if len(final_seat) != len(I):
        missing_diners = (
            set(I) - set(final_seat)
        )

        raise RuntimeError(
            "Not every diner received a seat. "
            f"Missing: {sorted(missing_diners)}"
        )

    if (
        len(set(final_seat.values()))
        != len(final_seat)
    ):
        raise RuntimeError(
            "More than one diner has been "
            "assigned to the same seat."
        )


    # OUTPUT
    output_rows = []

    for diner_id in I:
        diner = diner_by_id[diner_id]
        seat_id = final_seat[diner_id]
        seat = seat_by_id[seat_id]

        output_rows.append({
            "seat_id": seat_id,
            "seat_section": seat["section"],
            "seat_table": seat["table"],
            "seat_side": seat["side"],
            "seat_position": seat["position"],
            "stage1_section": (
                section_assignment[diner_id]
            ),
            "diner_id": diner_id,
            "category": diner["category"],
            "title": diner["title"],
            "first": diner["first"],
            "surname": diner["surname"],
            "full_name": diner["name"],
            "matric": (
                diner["matric"]
                if diner["matric"] is not None
                else ""
            ),
            "gender": diner["gender"],
            "notes": diner["notes"],
        })

    output_rows.sort(
        key=lambda row: (
            row["seat_section"],
            row["seat_table"],
            row["seat_side"],
            int(row["seat_position"]),
        )
    )

    fieldnames = [
        "seat_id",
        "seat_section",
        "seat_table",
        "seat_side",
        "seat_position",
        "stage1_section",
        "diner_id",
        "category",
        "title",
        "first",
        "surname",
        "full_name",
        "matric",
        "gender",
        "notes",
    ]

    with OUTPUT_CSV.open(
        "w",
        encoding="utf-8",
        newline="",
    ) as file:
        writer = csv.DictWriter(
            file,
            fieldnames=fieldnames,
        )

        writer.writeheader()
        writer.writerows(
            output_rows
        )

    print(
        "\nTwo-stage solution completed."
    )
    print(
        "Stage 1 objective:",
        stage1.attributes.objval,
    )
    print(
        "Stage 2 objectives:",
        stage2_objectives,
    )
    print(
        "Stage 2 total objective:",
        sum(stage2_objectives.values()),
    )
    print(
        "Final solution written to:",
        OUTPUT_CSV,
    )


if __name__ == "__main__":
    main()